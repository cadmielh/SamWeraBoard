"""Blocuri repetitive detectate automat dintr-un document fără etichete: mai multe persoane cu aceeași structură,
fie în aceeași frază („X, cetățenia…, CNP … si Y, cetățenia…, CNP …” — kind='inline'), fie în paragrafe separate
consecutive (kind='paragraph'). Alese ca „repeat”, devin un {{#ASOCIATI}}/{{#ADMINISTRATORI}} obișnuit, care
scalează la orice număr de persoane prin motorul deja existent — spre deosebire de poziții fixe (ASOCIAT_1,
ASOCIAT_2, …), care nu au loc pentru o a treia persoană dacă șablonul a fost scris doar pentru două."""
import io
import json
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import blanks  # noqa: E402
from doc_filler import fill_docx  # noqa: E402


def make_doc(lines: list[str]) -> bytes:
    d = Document()
    for t in lines:
        d.add_paragraph(t)
    b = io.BytesIO(); d.save(b)
    return b.getvalue()


IDENTITY = ("{name}, cetatean roman, nascut la data de {d}, în {loc}, jud.{jud}, identificat cu CI seria {ser} "
            "nr. {nr} eliberat de {em} la {la}, CNP {cnp}")


def two_person_sentence(lead_in: str = "constituit de asociatii ") -> str:
    a = IDENTITY.format(name="……", d="……", loc="……", jud="……", ser="……", nr="……", em="……", la="……", cnp="……")
    b = IDENTITY.format(name="……", d="……", loc="……", jud="……", ser="……", nr="……", em="……", la="……", cnp="……")
    return f"{lead_in}{a} si {b}."


def person(nume: str) -> dict:
    return {"NUME": nume, "PRENUME": "x", "CNP": "1900101123456", "DATA_NASTERII": "01.01.1990", "LOCUL_NASTERII": "Arad",
            "JUDET": "Arad", "SERIE_ACT": "AR", "NR_ACT": "123456", "EMISA_DE": "SPCLEP Arad", "VALABILA_DE_LA": "01.01.2020"}


# ── detectare ──────────────────────────────────────────────────────────────────
def test_doua_persoane_in_aceeasi_fraza_sunt_detectate_ca_grup_inline():
    raw = make_doc([two_person_sentence()])
    groups = blanks.detect_groups(blanks.analyze(raw))
    assert len(groups) == 1
    g = groups[0]
    assert g["kind"] == "inline" and g["role"] == "ASOCIAT" and g["count"] == 2
    assert "asociați" in g["label"]


def test_paragrafe_consecutive_cu_aceeasi_structura_sunt_detectate_ca_grup_paragraph():
    raw = make_doc(["1. " + IDENTITY.format(name="……", d="……", loc="……", jud="……", ser="……", nr="……", em="……", la="……", cnp="……"),
                    "2. " + IDENTITY.format(name="……", d="……", loc="……", jud="……", ser="……", nr="……", em="……", la="……", cnp="……")])
    groups = blanks.detect_groups(blanks.analyze(raw))
    assert len(groups) == 1 and groups[0]["kind"] == "paragraph" and groups[0]["count"] == 2


def test_o_singura_persoana_nu_formeaza_grup():
    raw = make_doc([IDENTITY.format(name="……", d="……", loc="……", jud="……", ser="……", nr="……", em="……", la="……", cnp="……") + "."])
    assert blanks.detect_groups(blanks.analyze(raw)) == []


def test_adresa_lipsa_la_o_persoana_nu_blocheaza_gruparea():
    """Caz real (ACTUL CONSTITUTIV): un asociat e scris „domiciliat în jud.……” (doar județ, fără localitate
    separată), celălalt „domiciliat în ……, jud. ……” (localitate + județ) — aceeași clauză, doar o mențiune mai
    scurtă a domiciliului. ADRESA e opțională la potrivirea formei tocmai pentru asta: fără relaxarea asta,
    gruparea se abținea și a treia persoană dispărea din acel paragraf (raportat de utilizator)."""
    a = IDENTITY.format(name="……", d="……", loc="……", jud="……", ser="……", nr="……", em="……", la="……", cnp="……")
    b_with_extra_address = ("……, cetatean roman, nascut la data de …… în ……, jud.……, domiciliat în ……, "
                            "identificat cu CI seria …… nr. …… eliberat de …… la ……, CNP ……")
    raw = make_doc([f"constituit de {a} si {b_with_extra_address}."])
    groups = blanks.detect_groups(blanks.analyze(raw))
    assert len(groups) == 1
    g = groups[0]
    assert g["kind"] == "inline" and g["role"] == "ASOCIAT" and g["count"] == 2
    # clauza-șablon e cea mai completă (cu ADRESA) — altfel domiciliul ar dispărea la toată lumea, nu doar la a
    found = blanks.analyze(raw)
    by_id = {b["id"]: b for b in found}
    template_fields = [by_id[i]["field"] for i in g["template_blank_ids"]]
    assert "ADRESA" in template_fields

    tpl = blanks.apply(raw, {}, {g["id"]: "repeat"})
    text_tpl = "\n".join(p.text for p in Document(io.BytesIO(tpl)).paragraphs)
    assert "{{#ASOCIATI}}" in text_tpl and "{{ADRESA}}" in text_tpl
    out = fill_docx(tpl, {}, groups={"ASOCIATI": [person(f"P{i}") | {"ADRESA": f"str. {i}"} for i in range(1, 4)]})
    text_out = "\n".join(p.text for p in Document(io.BytesIO(out)).paragraphs)
    assert text_out.count("cetatean roman") == 3            # toate cele 3 persoane apar, nu doar 2


def test_structuri_asimetrice_altfel_decat_adresa_nu_se_grupeaza_fortat():
    """Documentare a limitei relaxării de mai sus: o diferență de formă care NU e doar ADRESA (aici: CNP lipsă la
    o persoană) tot blochează gruparea automată — mai bine poziții fixe, protejate de avertismentul de
    capacitate, decât un șablon greșit."""
    a = IDENTITY.format(name="……", d="……", loc="……", jud="……", ser="……", nr="……", em="……", la="……", cnp="……")
    b_without_cnp = ("……, cetatean roman, nascut la data de …… în ……, jud.……, identificat cu CI seria …… "
                     "nr. …… eliberat de …… la ……")
    raw = make_doc([f"constituit de {a} si {b_without_cnp}."])
    assert blanks.detect_groups(blanks.analyze(raw)) == []


def test_administratori_si_asociati_sunt_grupuri_separate():
    raw = make_doc([two_person_sentence("asociatii "), two_person_sentence("administratorii ").replace("asociatii", "administratorii")])
    # al doilea paragraf trebuie recunoscut ca ADMINISTRATOR pentru a forma un grup propriu
    found = blanks.analyze(raw)
    groups = blanks.detect_groups(found)
    roles = {g["role"] for g in groups}
    assert "ASOCIAT" in roles


# ── aplicare + scalare reală (1, 3+ persoane) ──────────────────────────────────
def test_grup_inline_ales_repeat_scaleaza_la_orice_numar():
    raw = make_doc([two_person_sentence()])
    found = blanks.analyze(raw)
    gid = blanks.detect_groups(found)[0]["id"]
    tpl = blanks.apply(raw, {}, {gid: "repeat"})
    text_tpl = "\n".join(p.text for p in Document(io.BytesIO(tpl)).paragraphs)
    assert "{{#ASOCIATI}}" in text_tpl and "{{/ASOCIATI}}" in text_tpl
    assert "{{NUME}} {{PRENUME}}" in text_tpl and "{{CNP}}" in text_tpl

    for n in (1, 4):
        out = fill_docx(tpl, {}, groups={"ASOCIATI": [person(f"P{i}") for i in range(1, n + 1)]})
        text = "\n".join(p.text for p in Document(io.BytesIO(out)).paragraphs)
        assert "{{" not in text and "……" not in text
        assert all(f"P{i}" in text for i in range(1, n + 1))


def test_grup_paragraph_ales_repeat_pastreaza_numerotarea_index():
    raw = make_doc(["1. " + IDENTITY.format(name="……", d="……", loc="……", jud="……", ser="……", nr="……", em="……", la="……", cnp="……"),
                    "2. " + IDENTITY.format(name="……", d="……", loc="……", jud="……", ser="……", nr="……", em="……", la="……", cnp="……")])
    found = blanks.analyze(raw)
    gid = blanks.detect_groups(found)[0]["id"]
    tpl = blanks.apply(raw, {}, {gid: "repeat"})
    assert "{{INDEX}}." in "\n".join(p.text for p in Document(io.BytesIO(tpl)).paragraphs)

    out = fill_docx(tpl, {}, groups={"ASOCIATI": [person("P1"), person("P2"), person("P3")]})
    lines = [p.text for p in Document(io.BytesIO(out)).paragraphs if p.text.strip()]
    assert lines == [f"{i}. P{i} x, cetatean roman, nascut la data de 01.01.1990, în Arad, jud.Arad, "
                     f"identificat cu CI seria AR nr. 123456 eliberat de SPCLEP Arad la 01.01.2020, CNP 1900101123456"
                     for i in (1, 2, 3)]


def test_grup_ales_fixed_pastreaza_pozitiile_numerotate_de_dinainte():
    raw = make_doc([two_person_sentence()])
    found = blanks.analyze(raw)
    choices = {f["id"]: f["tag"] for f in found}       # utilizatorul acceptă propunerile implicite (numerotate)
    gid = blanks.detect_groups(found)[0]["id"]
    tpl = blanks.apply(raw, choices, {gid: "fixed"})
    text = "\n".join(p.text for p in Document(io.BytesIO(tpl)).paragraphs)
    assert "{{#ASOCIATI}}" not in text
    assert "{{ASOCIAT_1_NUME}} {{ASOCIAT_1_PRENUME}}" in text and "{{ASOCIAT_2_NUME}} {{ASOCIAT_2_PRENUME}}" in text


def test_fara_alegere_de_grup_ramane_ca_inainte_pozitii_fixe():
    raw = make_doc([two_person_sentence()])
    choices = {f["id"]: f["tag"] for f in blanks.analyze(raw)}
    tpl = blanks.apply(raw, choices)     # fără parametrul `groups` — comportament neschimbat
    text = "\n".join(p.text for p in Document(io.BytesIO(tpl)).paragraphs)
    assert "{{ASOCIAT_1_NUME}} {{ASOCIAT_1_PRENUME}}" in text


def test_grupul_inline_pastreaza_prefixul_si_sufixul_fix():
    raw = make_doc(["Constituit de asociatii " + two_person_sentence("")[0:-1] + ", acestia fiind de acord."])
    found = blanks.analyze(raw)
    groups = blanks.detect_groups(found)
    assert len(groups) == 1
    tpl = blanks.apply(raw, {}, {groups[0]["id"]: "repeat"})
    text = "\n".join(p.text for p in Document(io.BytesIO(tpl)).paragraphs)
    assert text.startswith("Constituit de asociatii") and text.rstrip().endswith("acestia fiind de acord.")


def test_blank_neutilizat_in_grup_primeste_alegerea_individuala():
    """Un loc liber din PREFIXUL/SUFIXUL paragrafului grupat (nu parte din nicio clauză) rămâne controlabil individual."""
    raw = make_doc(["Cu sediul in …… , constituit de " + two_person_sentence("")])
    found = blanks.analyze(raw)
    sediu = next(f for f in found if f["field"] is None or f["before"].strip().endswith("sediul in"))
    groups = blanks.detect_groups(found)
    assert len(groups) == 1
    gid = groups[0]["id"]
    tpl = blanks.apply(raw, {sediu["id"]: "{{SOCIETATE_SEDIU}}"}, {gid: "repeat"})
    text = "\n".join(p.text for p in Document(io.BytesIO(tpl)).paragraphs)
    assert "{{SOCIETATE_SEDIU}}" in text and "……" not in text.split("{{#ASOCIATI}}")[0]


# ── rută ────────────────────────────────────────────────────────────────────────
def test_ruta_blanks_intoarce_grupurile_si_apply_le_accepta(monkeypatch):
    import app as app_module
    monkeypatch.setattr(app_module, "_verify", lambda *a, **k: "uid")
    c = app_module.app.test_client()
    raw = make_doc([two_person_sentence()])

    r = c.post("/template/blanks", data={"template": (io.BytesIO(raw), "t.docx")}, content_type="multipart/form-data")
    body = r.get_json()
    assert len(body["groups"]) == 1
    gid = body["groups"][0]["id"]

    applied = c.post("/template/blanks/apply", data={
        "template": (io.BytesIO(raw), "t.docx"), "choices": "{}", "groups": json.dumps({str(gid): "repeat"}),
    }, content_type="multipart/form-data")
    assert applied.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(applied.data)).paragraphs)
    assert "{{#ASOCIATI}}" in text

    bad = c.post("/template/blanks/apply", data={
        "template": (io.BytesIO(raw), "t.docx"), "choices": "{}", "groups": "nu e json",
    }, content_type="multipart/form-data")
    assert bad.status_code == 400


# ── mențiuni de nume fără date de identitate: „…… și ……” → un singur câmp scalabil, nu poziții fixe ────────────
def test_doua_mentiuni_devin_un_singur_camp_lista_scalabila():
    raw = make_doc(["În calitate de asociati …… si …… au drepturile."])
    found = blanks.analyze(raw)
    assert len(found) == 1 and found[0]["scope"] == "company" and found[0]["field"] == "ASOCIATI_LISTA"


def test_trei_mentiuni_separate_prin_virgula_devin_tot_un_singur_camp():
    raw = make_doc(["Sunt asociati …… , …… si …… ."])
    found = blanks.analyze(raw)
    assert len(found) == 1 and found[0]["field"] == "ASOCIATI_LISTA"


def test_administratori_mentionati_devin_administratori_lista_nu_asociati_lista():
    raw = make_doc(["administratorii …… si …… vor semna actele."])
    found = blanks.analyze(raw)
    assert len(found) == 1 and found[0]["field"] == "ADMINISTRATORI_LISTA"


def test_o_singura_mentiune_ramane_individuala_nu_devine_lista():
    """O mențiune singură poate desemna o persoană anume (ex. administratorul nou numit) — nu se presupune
    că înseamnă „toți”, doar pentru că nu are date de identitate lângă ea."""
    raw = make_doc(["Reprezentantul legal al societății va fi asociatul …… . Reprezentarea este generală."])
    found = blanks.analyze(raw)
    assert len(found) == 1 and found[0]["field"] == "NUME_COMPLET" and found[0]["tag"] == "{{ASOCIAT_1_NUME}} {{ASOCIAT_1_PRENUME}}"


def test_lista_scalabila_functioneaza_la_1_2_si_3_persoane():
    raw = make_doc(["Sunt asociati …… , …… si …… ."])
    found = blanks.analyze(raw)
    tpl = blanks.apply(raw, {found[0]["id"]: found[0]["tag"]})
    text = Document(io.BytesIO(tpl)).paragraphs[0].text
    assert text == "Sunt asociati {{ASOCIATI_LISTA}} ."
    for lista in ("Ionescu Maria", "Ionescu Maria și Popescu Ion", "Ionescu Maria, Popescu Ion și Vasile Dan"):
        out = fill_docx(tpl, {"{{ASOCIATI_LISTA}}": lista})
        assert lista in Document(io.BytesIO(out)).paragraphs[0].text and "{{" not in Document(io.BytesIO(out)).paragraphs[0].text


def test_actul_constitutiv_real_recunoaste_lista_de_asociati():
    """Documentare a unui caz real: p31 („În calitate de asociati …… si …… au drepturile”) — devenea poziții
    fixe (2 locuri, fără loc pentru un al treilea asociat); acum devine ASOCIATI_LISTA."""
    raw = (Path(__file__).resolve().parent.parent / "fisiere_template" / "SabloaneAvocatTavi" / "anonimizate" / "ACTUL CONSTITUTIV.docx")
    if not raw.exists():
        return
    found = blanks.analyze(raw.read_bytes())
    lista = [f for f in found if f["field"] == "ASOCIATI_LISTA"]
    assert len(lista) == 1


# ── sediul defalcat (localitate + județ separat) vs. întreg ────────────────────────────────────────────────────
def test_sediul_defalcat_cu_judet_separat_alaturi_nu_dubleaza_judetul():
    raw = make_doc(["cu sediul în …… , jud. …… ."])
    found = blanks.analyze(raw)
    assert [f["field"] for f in found] == ["SOCIETATE_SEDIU_FARA_JUDET", "SOCIETATE_JUDET"]
    tpl = blanks.apply(raw, {f["id"]: f["tag"] for f in found})
    out = fill_docx(tpl, {"{{SOCIETATE_SEDIU_FARA_JUDET}}": "Timișoara, Str. Exemplu, nr. 1", "{{SOCIETATE_JUDET}}": "Timiș"})
    text = Document(io.BytesIO(out)).paragraphs[0].text
    assert text == "cu sediul în Timișoara, Str. Exemplu, nr. 1 , jud. Timiș ."
    assert text.count("jud.") == 1


def test_sediul_intreg_fara_judet_separat_ramane_campul_complet():
    raw = make_doc(["cu sediul în …… ."])
    found = blanks.analyze(raw)
    assert [f["field"] for f in found] == ["SOCIETATE_SEDIU"]


def test_actul_constitutiv_real_are_sediul_intreg():
    """Sediul e scris ca un singur loc liber („îşi are sediul în ………….”), fără județ defalcat separat alături —
    rămâne SOCIETATE_SEDIU complet. Un SOCIETATE_JUDET apare separat, la p17 (banca „de pe teritoriul jud.……”),
    fără legătură cu sediul."""
    raw = Path(__file__).resolve().parent.parent / "fisiere_template" / "SabloaneAvocatTavi" / "anonimizate" / "ACTUL CONSTITUTIV.docx"
    if not raw.exists():
        return
    found = blanks.analyze(raw.read_bytes())
    fields = [f["field"] for f in found if f["field"] in ("SOCIETATE_SEDIU", "SOCIETATE_SEDIU_FARA_JUDET", "SOCIETATE_JUDET")]
    assert fields == ["SOCIETATE_SEDIU", "SOCIETATE_JUDET"]


# ── font ─────────────────────────────────────────────────────────────────────
def test_paragrafele_generate_pastreaza_fontul_sablonului_nu_cad_pe_cel_implicit():
    """Bug raportat de utilizator: textul completat ieșea mereu cu Times New Roman, indiferent de fontul
    din șablon (ex. Arial), pentru paragrafele generate la transformarea unui bloc repetitiv."""
    from docx.oxml.ns import qn

    d = Document()
    p = d.add_paragraph()
    run = p.add_run(two_person_sentence())
    run.font.name = "Arial"
    b = io.BytesIO(); d.save(b)
    raw = b.getvalue()

    groups = blanks.detect_groups(blanks.analyze(raw))
    tpl = blanks.apply(raw, {}, {groups[0]["id"]: "repeat"})
    out_doc = Document(io.BytesIO(tpl))

    def font_of(paragraph):
        for r in paragraph.runs:
            if not r.text:
                continue
            rpr = r._r.find(qn("w:rPr"))
            rfonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
            return rfonts.get(qn("w:ascii")) if rfonts is not None else None
        return None

    fonts = [font_of(p) for p in out_doc.paragraphs if p.text.strip()]
    assert fonts and all(f == "Arial" for f in fonts)


# ── CAEN (obiectul de activitate), needitat de la firma-exemplu ───────────────
def test_domeniul_si_activitatea_principala_needitate_devin_caen_domeniu_si_caen_1():
    """„Domeniul principal de activitate este: 953 Repararea ...” / „Activitatea principală este: 9531
    Repararea ...” — needitate de la firma-exemplu (fără „……”) — devin {{CAEN_DOMENIU}} (grupa, 3 cifre) și
    {{CAEN_1}} (clasa, 4 cifre) — două câmpuri distincte, aplicația derivă automat grupa din clasă."""
    raw = make_doc([
        "Art.7.- Domeniul principal de activitate este: 953 Repararea și întreținerea autovehiculelor ; -",
        "Activitatea principală este: 9531 Repararea și întreținerea autovehiculelor  ---------",
    ])
    found = blanks.analyze(raw)
    assert [f["field"] for f in found] == ["CAEN_DOMENIU", "CAEN_1"]
    assert all(f["confidence"] == "high" for f in found)

    tpl = blanks.apply(raw, {f["id"]: f["tag"] for f in found})
    out = fill_docx(tpl, {"{{CAEN_DOMENIU}}": "953 - Repararea și întreținerea autovehiculelor",
                          "{{CAEN_1}}": "9531 - Repararea și întreținerea autovehiculelor"})
    texts = [p.text for p in Document(io.BytesIO(out)).paragraphs]
    assert texts[0] == "Art.7.- Domeniul principal de activitate este: 953 - Repararea și întreținerea autovehiculelor ; -"
    assert texts[1] == "Activitatea principală este: 9531 - Repararea și întreținerea autovehiculelor  ---------"


def test_lista_activitati_secundare_needitata_devine_grup_repetitiv_caen():
    """„Societatea va mai desfăsura si următoarele activităti:” urmat de coduri CAEN, câte unul pe rând, fără
    nicio frază proprie — needitate de la firma-exemplu — devine {{#CAEN_SECUNDARE}}, scalabil la orice număr."""
    raw = make_doc([
        "Societatea va mai desfăsura si următoarele activităti:",
        "4672 Comerţ cu ridicata al pieselor şi accesoriilor pentru autovehicule",
        "4690 Comerţ cu ridicata nespecializat",
        "4778 Comerţ cu amănuntul al altor bunuri noi",
    ])
    found = blanks.analyze(raw)
    caen_lines = [f for f in found if f["field"] == "CAEN"]
    assert len(caen_lines) == 3 and all(f["role"] == "CAEN" for f in caen_lines)

    groups = blanks.detect_groups(found)
    assert len(groups) == 1
    g = groups[0]
    assert g["kind"] == "paragraph" and g["role"] == "CAEN" and g["count"] == 3

    tpl = blanks.apply(raw, {}, {g["id"]: "repeat"})
    text_tpl = "\n".join(p.text for p in Document(io.BytesIO(tpl)).paragraphs)
    assert "{{#CAEN_SECUNDARE}}" in text_tpl and "{{CAEN}}" in text_tpl and "{{/CAEN_SECUNDARE}}" in text_tpl
    # declanșatorul rămâne text fix, needitat
    assert "următoarele activităti" in text_tpl

    for n in (1, 5):
        out = fill_docx(tpl, {}, groups={"CAEN_SECUNDARE": [{"CAEN": f"C{i}"} for i in range(1, n + 1)]})
        texts = [p.text for p in Document(io.BytesIO(out)).paragraphs if p.text.strip()]
        assert texts[1:] == [f"C{i}" for i in range(1, n + 1)]


def test_actul_constitutiv_real_recunoaste_sectiunea_caen_completa():
    """Documentul are acum două forme, ambele recunoscute: „Domeniul principal de activitate este: ……” (loc
    liber simplu, fără indiciu → {{CAEN_DOMENIU}}, grupa) și „Activitatea principală este: …(CAEN PRINCIPAL)……”
    / „…(CAEN SECUNDARE)……….” (cu indiciu explicit din paranteze, scris de utilizator) — vezi _resolve_hint."""
    raw = Path(__file__).resolve().parent.parent / "fisiere_template" / "SabloaneAvocatTavi" / "anonimizate" / "ACTUL CONSTITUTIV.docx"
    if not raw.exists():
        return
    data = raw.read_bytes()
    found = blanks.analyze(data)
    assert sum(f["field"] == "CAEN_DOMENIU" for f in found) == 1              # domeniul (grupa, fără indiciu)
    assert sum(f["field"] == "CAEN_1" for f in found) == 1                    # activitatea principală (indiciu)
    caen_secundare = [f for f in found if f["field"] == "CAEN"]
    assert len(caen_secundare) == 1                                          # un singur indiciu „(CAEN SECUNDARE)”, care scalează oricum

    groups = blanks.detect_groups(found)
    caen_groups = [g for g in groups if g["role"] == "CAEN"]
    assert len(caen_groups) == 1 and caen_groups[0]["count"] == 1

    gids = {g["id"]: "repeat" for g in groups}
    grouped_ids = {i for g in groups for i in g["blank_ids"]}
    choices = {f["id"]: f["tag"] for f in found if f["id"] not in grouped_ids and f["scope"] == "company"}
    tpl = blanks.apply(data, choices, gids)
    out = fill_docx(tpl, {"{{CAEN_1}}": "7022 - Activități de consultanță"},
                    groups={"ASOCIATI": [], "ADMINISTRATORI": [],
                            "CAEN_SECUNDARE": [{"CAEN": f"COD{i}"} for i in range(1, 4)]})
    texts = [p.text for p in Document(io.BytesIO(out)).paragraphs]
    assert texts.count("COD1") == 1 and texts.count("COD2") == 1 and texts.count("COD3") == 1
    assert texts.count("7022 - Activități de consultanță") == 0              # nu apare singur — mereu în frază
    assert any("7022 - Activități de consultanță" in t for t in texts)


# ── indiciu explicit din paranteze, lipit de locul liber ───────────────────────
def test_indiciu_intre_doua_locuri_libere_devine_o_singura_eticheta():
    """„…(CAEN PRINCIPAL)……” — indiciul scris de utilizator, lipit de locul liber, are prioritate maximă;
    cele două locuri libere + indiciul devin O SINGURĂ etichetă, nu rămân separate."""
    raw = make_doc(["Activitatea principală este: …(CAEN PRINCIPAL)……."])
    found = blanks.analyze(raw)
    assert len(found) == 1
    f = found[0]
    assert f["scope"] == "company" and f["field"] == "CAEN_1" and f["confidence"] == "high"

    tpl = blanks.apply(raw, {f["id"]: f["tag"]})
    out = fill_docx(tpl, {"{{CAEN_1}}": "6201 - Activități de realizare a soft-ului la comandă"})
    text = Document(io.BytesIO(out)).paragraphs[0].text
    assert text == "Activitatea principală este: 6201 - Activități de realizare a soft-ului la comandă."
    assert "(" not in text and ")" not in text          # indiciul nu rămâne vizibil în documentul generat


def test_indiciu_caen_secundare_singur_devine_grup_repetitiv_de_1():
    """„…(CAEN SECUNDARE)……….” — un singur indiciu, fără listă de exemple alături — tot devine bloc repetitiv
    {{#CAEN_SECUNDARE}}, scalabil la câte activități secundare are clientul, nu doar la una."""
    raw = make_doc(["Societatea va mai desfășura și următoarele activități:", "…(CAEN SECUNDARE)……….",])
    found = blanks.analyze(raw)
    caen = [f for f in found if f["field"] == "CAEN"]
    assert len(caen) == 1 and caen[0]["role"] == "CAEN"

    groups = blanks.detect_groups(found)
    assert len(groups) == 1 and groups[0]["role"] == "CAEN" and groups[0]["count"] == 1
    tpl = blanks.apply(raw, {}, {groups[0]["id"]: "repeat"})
    out = fill_docx(tpl, {}, groups={"CAEN_SECUNDARE": [{"CAEN": f"C{i}"} for i in range(1, 4)]})
    texts = [p.text for p in Document(io.BytesIO(out)).paragraphs if p.text.strip()]
    assert texts[1:] == ["C1", "C2", "C3", "."]     # punctul final al frazei rămâne text fix, o singură dată, după listă


def test_indiciu_necunoscut_devine_camp_manual_etichetat_din_indiciu():
    """Un indiciu care nu corespunde niciunui câmp cunoscut (ex. „(Suma penalitate)”) devine automat câmp
    manual, cu indiciul ca etichetă — mai bun decât ghicirea din ultimele cuvinte dinainte de loc liber."""
    raw = make_doc(["Suma datorată este de …(Suma penalitate)…… lei."])
    found = blanks.analyze(raw)
    assert len(found) == 1
    f = found[0]
    assert f["scope"] == "manual" and f["label"] == "Suma penalitate" and f["confidence"] == "high"
    assert f["tag"] == "{{CAMP_SUMA_PENALITATE}}"


def test_indiciu_cu_numele_intern_al_campului_functioneaza_direct():
    """Utilizatorul poate scrie direct numele intern al câmpului în paranteză (ex. „(SOCIETATE_SEDIU)”), nu
    doar aliasurile cunoscute — util când nicio formulare din alias-uri nu i se potrivește."""
    raw = make_doc(["Sediul social este în …(SOCIETATE_SEDIU)…… ."])
    found = blanks.analyze(raw)
    assert found[0]["scope"] == "company" and found[0]["field"] == "SOCIETATE_SEDIU"


def test_domeniul_si_activitatea_principala_fara_indiciu_raman_recunoscute():
    """Fără niciun indiciu, doar loc liber simplu după formulările uzuale — devin {{CAEN_DOMENIU}} / {{CAEN_1}}, din context."""
    raw = make_doc([
        "Domeniul principal de activitate este: …………………",
        "Activitatea principală este: …….",
    ])
    found = blanks.analyze(raw)
    assert [f["field"] for f in found] == ["CAEN_DOMENIU", "CAEN_1"]
    assert all(f["confidence"] == "high" for f in found)


# ── o singură clauză, dar cu marcaj explicit de plural („asociatul/asociații ……”) ──────────────────────────────
def test_o_singura_clauza_cu_marcaj_asociatul_asociatii_devine_grup_scalabil():
    """Cerință directă a utilizatorului: „acolo unde am pus asociatul/asociații ... și o singură suită de
    câmpuri, mă aștept să multiplici acel paragraf de câte ori e cazul”. O singură persoană-exemplu, dar
    marcajul de plural alături, devine grup — la fel ca la 2+ clauze, scalează la orice număr."""
    a = IDENTITY.format(name="……", d="……", loc="……", jud="……", ser="……", nr="……", em="……", la="……", cnp="……")
    raw = make_doc([f"constituită de asociatul/asociații {a}."])
    found = blanks.analyze(raw)
    groups = blanks.detect_groups(found, raw)
    assert len(groups) == 1
    g = groups[0]
    assert g["kind"] == "inline" and g["role"] == "ASOCIAT" and g["count"] == 1

    tpl = blanks.apply(raw, {}, {g["id"]: "repeat"})
    text_tpl = "\n".join(p.text for p in Document(io.BytesIO(tpl)).paragraphs)
    assert "{{#ASOCIATI}}" in text_tpl and "{{NUME}} {{PRENUME}}" in text_tpl
    for n in (1, 3):
        out = fill_docx(tpl, {}, groups={"ASOCIATI": [person(f"P{i}") for i in range(1, n + 1)]})
        text = "\n".join(p.text for p in Document(io.BytesIO(out)).paragraphs)
        assert all(f"P{i}" in text for i in range(1, n + 1))


def test_marcajul_de_plural_intr_un_paragraf_separat_dinainte_tot_functioneaza():
    """Caz real (ACTUL CONSTITUTIV): „...asociatul/asociații” rămâne singur, ca ultim rând al unui paragraf
    Word, iar persoana cu datele complete începe într-un paragraf SEPARAT, imediat următor — tot trebuie
    recunoscut, nu doar când marcajul e în același paragraf cu persoana."""
    a = IDENTITY.format(name="……", d="……", loc="……", jud="……", ser="……", nr="……", em="……", la="……", cnp="……")
    raw = make_doc(["constituită de asociatul/asociații", a + "."])
    found = blanks.analyze(raw)
    groups = blanks.detect_groups(found, raw)
    assert len(groups) == 1 and groups[0]["role"] == "ASOCIAT" and groups[0]["paragraph"] == 1


def test_administrator_administratori_ca_marcaj_de_plural():
    a = IDENTITY.format(name="……", d="……", loc="……", jud="……", ser="……", nr="……", em="……", la="……", cnp="……")
    raw = make_doc([f"administrarea e îndeplinită de administratorul/administratorii {a}."])
    found = blanks.analyze(raw)
    groups = blanks.detect_groups(found, raw)
    assert len(groups) == 1 and groups[0]["role"] == "ADMINISTRATOR"


def test_fara_marcaj_de_plural_o_singura_clauza_nu_devine_grup():
    """Păstrează comportamentul actual: fără „X/Y” alături, o singură persoană rămâne poziție simplă (1),
    nu se oferă nicio grupare — exact cerința utilizatorului („dar păstrează și cum faci acum”)."""
    a = IDENTITY.format(name="……", d="……", loc="……", jud="……", ser="……", nr="……", em="……", la="……", cnp="……")
    raw = make_doc([f"constituită de asociatul {a}."])
    found = blanks.analyze(raw)
    assert blanks.detect_groups(found, raw) == []


def test_fara_docx_bytes_grupul_cu_o_clauza_nu_se_ofera_dar_restul_functioneaza():
    """Compatibilitate: apelul vechi detect_groups(found), fără al doilea argument, nu se strică — doar nu
    beneficiază de grupul cu o singură clauză (are nevoie de textul paragrafului dinainte)."""
    raw = make_doc([two_person_sentence()])
    found = blanks.analyze(raw)
    assert blanks.detect_groups(found) == blanks.detect_groups(found, raw)   # 2+ clauze nu depind de docx_bytes
