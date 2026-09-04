"""Teste unitare pentru doc_filler.py — biblioteca de clauze opționale
({{#CLAUZE}}...{{/CLAUZE}}) și non-regresia motorului existent (placeholdere
plate + blocuri repetitive {{#TAG}}...{{/TAG}})."""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document

import doc_filler as df


def _build_docx(paragraphs: list[str]) -> bytes:
    """Construiește un .docx minimal, un paragraf per element din listă."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _read_paragraphs(docx_bytes: bytes) -> list[str]:
    doc = Document(io.BytesIO(docx_bytes))
    return ["".join(r.text for r in p.runs) for p in doc.paragraphs]


CLAUSE_TEMPLATE = [
    "Titlu document",
    "{{#CLAUZE}}",
    "Denumire: CLAUZA A",
    "Art. {{ART_NR}} Text clauza A cu {{CAMP_A}}.",
    "Denumire: CLAUZA B",
    "Art. {{ART_NR}} Text clauza B cu {{CAMP_B}}.",
    "Denumire: CLAUZA C",
    "Art. {{ART_NR}} Text clauza C cu {{CAMP_C}}.",
    "{{/CLAUZE}}",
    "Semnătură",
]


# ── list_clauses_in_docx ────────────────────────────────────────────────────

def test_list_clauses_detects_all_three_in_order():
    clauses = df.list_clauses_in_docx(_build_docx(CLAUSE_TEMPLATE))
    assert [c["tag"] for c in clauses] == ["CLAUZA_A", "CLAUZA_B", "CLAUZA_C"]
    assert clauses[0]["label"] == "CLAUZA A"
    assert clauses[0]["placeholders"] == ["{{CAMP_A}}"]


def test_list_clauses_excludes_art_nr():
    clauses = df.list_clauses_in_docx(_build_docx(CLAUSE_TEMPLATE))
    for c in clauses:
        assert "{{ART_NR}}" not in c["placeholders"]


def test_list_clauses_empty_without_markers():
    assert df.list_clauses_in_docx(_build_docx(["Text simplu", "{{NUME}}"])) == []


# ── fill_docx cu selected_clauses ───────────────────────────────────────────

def test_fill_docx_keeps_only_selected_clauses_in_template_order():
    replacements = {"{{CAMP_A}}": "A", "{{CAMP_B}}": "B", "{{CAMP_C}}": "C"}
    out = df.fill_docx(_build_docx(CLAUSE_TEMPLATE), replacements, selected_clauses=["CLAUZA_C", "CLAUZA_A"])
    joined = "\n".join(_read_paragraphs(out))
    assert "clauza A" in joined
    assert "clauza C" in joined
    assert "clauza B" not in joined
    assert "Denumire:" not in joined
    assert "{{#CLAUZE}}" not in joined and "{{/CLAUZE}}" not in joined


def test_fill_docx_numbers_art_nr_sequentially_over_kept_clauses_only():
    # Selectate în ordine "greșită" (C înaintea A) — ordinea din document tot câștigă.
    replacements = {"{{CAMP_A}}": "A", "{{CAMP_C}}": "C"}
    out = df.fill_docx(_build_docx(CLAUSE_TEMPLATE), replacements, selected_clauses=["CLAUZA_C", "CLAUZA_A"])
    texts = [t for t in _read_paragraphs(out) if t.strip()]
    assert "Art. 1 Text clauza A cu A." in texts
    assert "Art. 2 Text clauza C cu C." in texts


def test_fill_docx_single_clause_selected():
    out = df.fill_docx(_build_docx(CLAUSE_TEMPLATE), {"{{CAMP_B}}": "B"}, selected_clauses=["CLAUZA_B"])
    texts = [t for t in _read_paragraphs(out) if t.strip()]
    assert "Art. 1 Text clauza B cu B." in texts
    assert not any("clauza A" in t or "clauza C" in t for t in texts)


def test_fill_docx_no_clauses_selected_removes_all_articles():
    out = df.fill_docx(_build_docx(CLAUSE_TEMPLATE), {}, selected_clauses=[])
    texts = [t for t in _read_paragraphs(out) if t.strip()]
    assert texts == ["Titlu document", "Semnătură"]


# ── Non-regresie ────────────────────────────────────────────────────────────

def test_fill_docx_without_clauze_markers_behaves_as_before():
    """Un document fără {{#CLAUZE}} nu e afectat de selected_clauses — motorul
    existent (placeholdere plate) rămâne neschimbat, indiferent ce se pasează."""
    template = _build_docx(["Bună, {{NUME}}!"])
    out_without = df.fill_docx(template, {"{{NUME}}": "Ion"})
    out_with = df.fill_docx(template, {"{{NUME}}": "Ion"}, selected_clauses=["ORICE"])
    assert out_without == out_with
    assert "Bună, Ion!" in _read_paragraphs(out_without)


def test_fill_docx_repeat_blocks_still_work_alongside_clause_library():
    """{{#ASOCIATI}} (bloc repetitiv existent) funcționează neschimbat chiar și
    într-un document care are și o secțiune {{#CLAUZE}}."""
    template = _build_docx([
        "{{#ASOCIATI}}",
        "{{INDEX}}. {{NUME}}",
        "{{/ASOCIATI}}",
        "{{#CLAUZE}}",
        "Denumire: CLAUZA A",
        "Art. {{ART_NR}} {{CAMP_A}}",
        "{{/CLAUZE}}",
    ])
    groups = {"ASOCIATI": [{"NUME": "Popescu"}, {"NUME": "Ionescu"}]}
    out = df.fill_docx(template, {"{{CAMP_A}}": "text"}, groups=groups, selected_clauses=["CLAUZA_A"])
    texts = [t for t in _read_paragraphs(out) if t.strip()]
    assert "1. Popescu" in texts
    assert "2. Ionescu" in texts
    assert "Art. 1 text" in texts


# ── Tabele imbricate + repetare rânduri de tabel ───────────────────────────

def _build_docx_with_table(rows: list[list[str]]) -> bytes:
    """.docx cu un singur tabel top-level, un rând per element din `rows`,
    fiecare rând fiind o listă de texte de celulă (o singură celulă -> rândul
    e "merge-uit" pe toată lățimea, ca rândurile-marcaj din declarația ONRC)."""
    doc = Document()
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    for ri, cells in enumerate(rows):
        for ci, text in enumerate(cells):
            table.cell(ri, ci).text = text
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _build_docx_with_nested_table(outer_cell_text: str, inner_rows: list[list[str]]) -> bytes:
    """.docx cu un tabel top-level de un rând/o celulă, care conține la rândul
    ei un tabel imbricat — reproduce structura reală a declarației ONRC unde
    tabelele 3.1/3.2/3.3 stau imbricate în celula secțiunii respective."""
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    outer_cell = outer.cell(0, 0)
    outer_cell.paragraphs[0].add_run(outer_cell_text)
    ncols = max(len(r) for r in inner_rows)
    inner = outer_cell.add_table(rows=len(inner_rows), cols=ncols)
    for ri, cells in enumerate(inner_rows):
        for ci, text in enumerate(cells):
            inner.cell(ri, ci).text = text
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _read_table_rows(docx_bytes: bytes, table_index: int = 0) -> list[list[str]]:
    doc = Document(io.BytesIO(docx_bytes))
    table = doc.tables[table_index]
    return [[cell.text for cell in df._row_distinct_cells(row)] for row in table.rows]


def test_expand_repeat_table_rows_clones_once_per_item():
    template = _build_docx_with_table([
        ["{{#CAEN_SEDIU}}"],
        ["{{CAEN}} - {{CAEN_DESC}}"],
        ["{{/CAEN_SEDIU}}"],
    ])
    row_groups = {"CAEN_SEDIU": [{"CAEN": "6201", "CAEN_DESC": "Activități IT"}, {"CAEN": "6202", "CAEN_DESC": "Consultanță IT"}]}
    out = df.fill_docx(template, {}, row_groups=row_groups)
    rows = _read_table_rows(out)
    assert rows == [["6201 - Activități IT"], ["6202 - Consultanță IT"]]


def test_expand_repeat_table_rows_numbers_index_per_item():
    template = _build_docx_with_table([
        ["{{#SEDII}}"],
        ["{{INDEX}}. {{ADRESA}}"],
        ["{{/SEDII}}"],
    ])
    row_groups = {"SEDII": [{"ADRESA": "Str. A"}, {"ADRESA": "Str. B"}]}
    out = df.fill_docx(template, {}, row_groups=row_groups)
    rows = _read_table_rows(out)
    assert rows == [["1. Str. A"], ["2. Str. B"]]


def test_expand_repeat_table_rows_empty_group_removes_template_row_entirely():
    template = _build_docx_with_table([
        ["Header"],
        ["{{#CAEN_TERTI}}"],
        ["{{CAEN}}"],
        ["{{/CAEN_TERTI}}"],
        ["Footer"],
    ])
    out = df.fill_docx(template, {}, row_groups={"CAEN_TERTI": []})
    rows = _read_table_rows(out)
    assert rows == [["Header"], ["Footer"]]


def test_expand_repeat_table_rows_no_matching_group_leaves_rows_untouched():
    template = _build_docx_with_table([
        ["{{#CAEN_SEDIU}}"],
        ["{{CAEN}}"],
        ["{{/CAEN_SEDIU}}"],
    ])
    out = df.fill_docx(template, {}, row_groups=None)
    rows = _read_table_rows(out)
    assert rows == [["{{#CAEN_SEDIU}}"], ["{{CAEN}}"], ["{{/CAEN_SEDIU}}"]]


def test_expand_repeat_table_rows_works_inside_nested_table():
    """Reproduce structura reală a declarației ONRC — tabelul cu rânduri
    repetitive stă imbricat într-o celulă a unui tabel exterior, nu la
    nivelul documentului."""
    template = _build_docx_with_nested_table("3.1 SEDIU SOCIAL", [
        ["{{#CAEN_SEDIU}}"],
        ["{{CAEN}}"],
        ["{{/CAEN_SEDIU}}"],
    ])
    row_groups = {"CAEN_SEDIU": [{"CAEN": "6201"}, {"CAEN": "6202"}, {"CAEN": "6203"}]}
    out = df.fill_docx(template, {}, row_groups=row_groups)
    doc = Document(io.BytesIO(out))
    outer_cell = doc.tables[0].cell(0, 0)
    nested = outer_cell.tables[0]
    rows = [[cell.text for cell in df._row_distinct_cells(row)] for row in nested.rows]
    assert rows == [["6201"], ["6202"], ["6203"]]


def test_flat_replacement_reaches_nested_table_cells():
    """Substituția {{PLACEHOLDER}} plată trebuie să ajungă și în celulele unui
    tabel imbricat, nu doar la nivelul documentului/tabelului exterior."""
    template = _build_docx_with_nested_table("Antet", [["{{NUME_FIRMA}}"]])
    out = df.fill_docx(template, {"{{NUME_FIRMA}}": "ACME SRL"})
    doc = Document(io.BytesIO(out))
    nested = doc.tables[0].cell(0, 0).tables[0]
    assert nested.cell(0, 0).text == "ACME SRL"


def test_list_placeholders_in_docx_finds_placeholders_in_nested_table():
    template = _build_docx_with_nested_table("Antet", [["{{NUME_FIRMA}}"]])
    assert df.list_placeholders_in_docx(template) == ["{{NUME_FIRMA}}"]


def test_fill_docx_nested_repeat_group_inside_a_clause():
    """Un bloc {{#TAG}} imbricat în interiorul unei clauze (ex. structura
    rezultată după cesiune) e expandat corect înainte de filtrarea clauzelor."""
    template = _build_docx([
        "{{#CLAUZE}}",
        "Denumire: CESIUNE",
        "Art. {{ART_NR}} Structura rezultată:",
        "{{#STRUCTURA}}",
        "{{NUME}} - {{PROCENT}}%",
        "{{/STRUCTURA}}",
        "{{/CLAUZE}}",
    ])
    groups = {"STRUCTURA": [{"NUME": "Ionescu", "PROCENT": "100"}]}
    out = df.fill_docx(template, {}, groups=groups, selected_clauses=["CESIUNE"])
    texts = [t for t in _read_paragraphs(out) if t.strip()]
    assert "Ionescu - 100%" in texts
