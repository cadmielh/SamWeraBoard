"""Șabloane fără etichete: locurile libere se recunosc din context, se aplică după confirmare și documentul se completează normal."""
import io
import json
import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import blanks  # noqa: E402
from doc_filler import fill_docx  # noqa: E402

LINES = [
    "CONTRACT DE COMODAT",
    "……",
    "Art. 1. PARTILE",
    "……, domiciliată în ……, jud. ……, identificată cu CI seria …… nr. …… eliberată de …… la ……, CNP ……, in calitate de COMODANT si",
    "……, cu sediul in ……, jud. ……, C.U.I.:……, reprezentată de administrator ……, in calitate de COMODATAR",
    "Art. 3. Contractul se incheie pe o durată nelimitată, incepand cu ……, data la care bunul va fi predat.",
    "Prezentul contract s-a incheiat astazi …… in 3 exemplare.",
    "Capitalul social subscris este în valoare de …… lei, împărțit în …… părți sociale.",
    "Semnătura ________________",
]


def make_doc(lines=LINES) -> bytes:
    d = Document()
    for t in lines:
        d.add_paragraph(t)
    b = io.BytesIO(); d.save(b)
    return b.getvalue()


def by_tag(found):
    return [f["tag"] for f in found]


def test_recunoaste_campurile_din_context():
    found = blanks.analyze(make_doc())
    tags = by_tag(found)
    assert "{{SOCIETATE_DENUMIRE}}" in tags[:1]                      # denumirea firmei, singură sub titlu
    assert "{{ASOCIAT_1_NUME}} {{ASOCIAT_1_PRENUME}}" in tags         # persoana definită (nume urmat de date de identitate)
    for f in ("ADRESA", "JUDET", "SERIE_ACT", "NR_ACT", "EMISA_DE", "VALABILA_DE_LA", "CNP"):
        assert f"{{{{ASOCIAT_1_{f}}}}}" in tags, f
    # sediul e defalcat în document ("cu sediul in ……, jud. ……"): SOCIETATE_SEDIU_FARA_JUDET, nu SOCIETATE_SEDIU
    # (altfel județul ar apărea de două ori — o dată din adresă, o dată din locul liber separat)
    for t in ("{{SOCIETATE_SEDIU_FARA_JUDET}}", "{{SOCIETATE_JUDET}}", "{{SOCIETATE_CIF}}", "{{ADMINISTRATOR_1_NUME}} {{ADMINISTRATOR_1_PRENUME}}",
              "{{CAPITAL_SOCIAL_TOTAL}}", "{{PARTI_SOCIALE_TOTALE}}", "{{DATA_AZI}}"):
        assert t in tags, t


def test_ce_nu_se_recunoaste_devine_camp_manual_cu_eticheta_din_text():
    found = blanks.analyze(make_doc())
    manual = [f for f in found if f["scope"] == "manual"]
    assert len(manual) == 1 and manual[0]["tag"].startswith("{{CAMP_") and "INCEPAND" in manual[0]["tag"]
    assert manual[0]["confidence"] == "low"


def test_liniile_de_semnatura_nu_sunt_locuri_de_completat():
    found = blanks.analyze(make_doc())
    assert all("________" not in (f["before"] + f["after"]) or f["scope"] != "manual" for f in found)
    assert not any(f["paragraph"] == len(LINES) - 1 for f in found)


def test_persoana_definita_complet_ramane_numerotata_dar_mentiunile_devin_lista_scalabila():
    """Persoana 1 (nume + CNP + adresă etc.) rămâne individuală (poziție numerotată) — se poate lega de restul
    datelor ei. O simplă mențiune de nume, „…… si ……”, nu are loc pentru o a treia persoană dacă mai apare una:
    devine ASOCIATI_LISTA, un câmp care se adaptează singur la orice număr de asociați (vezi lib/placeholders.ts,
    joinNames)."""
    lines = ["Actul este constituit de asociatul ……, cetătean roman, născut la data de …… în ……, domiciliat în ……, CNP …… si ……, cetătean roman.",
             "Beneficiarii reali sunt asociatii, …… si …… .", "În calitate de asociati …… si …… au drepturile."]
    found = blanks.analyze(make_doc(lines))
    p1 = [f["tag"] for f in found if f["paragraph"] == 0 and f["field"] == "NUME_COMPLET"]
    assert p1 == ["{{ASOCIAT_1_NUME}} {{ASOCIAT_1_PRENUME}}", "{{ASOCIAT_2_NUME}} {{ASOCIAT_2_PRENUME}}"]
    for p in (1, 2):
        assert [f["tag"] for f in found if f["paragraph"] == p] == ["{{ASOCIATI_LISTA}}"]


def test_aplicarea_inlocuieste_doar_ce_e_ales_si_pastreaza_restul():
    raw = make_doc()
    found = blanks.analyze(raw)
    cnp = next(f for f in found if f["field"] == "CNP")
    out = blanks.apply(raw, {cnp["id"]: cnp["tag"]})
    text = "\n".join(p.text for p in Document(io.BytesIO(out)).paragraphs)
    assert "CNP {{ASOCIAT_1_CNP}}," in text
    assert text.count("……") == sum(f["end"] - f["start"] > 0 for f in found) - 1     # restul rămâne neatins


def test_eticheta_trimisa_de_client_e_validata():
    assert blanks.valid_tag("{{ASOCIAT_1_NUME}} {{ASOCIAT_1_PRENUME}}")
    assert blanks.valid_tag("{{CAMP_NR_HOTARARE}}")
    for bad in ("<script>", "{{a b}}", "text liber", "", None, "{{X}}" * 5, "{{lower}}"):
        assert not blanks.valid_tag(bad), bad


def test_flux_complet_locuri_libere_apoi_completare_cu_datele_clientului():
    raw = make_doc()
    found = blanks.analyze(raw)
    template = blanks.apply(raw, {f["id"]: f["tag"] for f in found})
    values = {
        "{{SOCIETATE_DENUMIRE}}": "EXEMPLU SRL", "{{SOCIETATE_SEDIU_FARA_JUDET}}": "Timișoara, Str. Exemplu nr. 1", "{{SOCIETATE_JUDET}}": "Timiș",
        "{{SOCIETATE_CIF}}": "12345678", "{{CAPITAL_SOCIAL_TOTAL}}": "200", "{{PARTI_SOCIALE_TOTALE}}": "20", "{{DATA_AZI}}": "01.01.2026",
        "{{ASOCIAT_1_NUME}}": "Ionescu", "{{ASOCIAT_1_PRENUME}}": "Maria", "{{ASOCIAT_1_ADRESA}}": "Cluj, str. Test 2", "{{ASOCIAT_1_JUDET}}": "Cluj",
        "{{ASOCIAT_1_SERIE_ACT}}": "CJ", "{{ASOCIAT_1_NR_ACT}}": "123456", "{{ASOCIAT_1_EMISA_DE}}": "SPCLEP Cluj", "{{ASOCIAT_1_VALABILA_DE_LA}}": "01.02.2020",
        "{{ASOCIAT_1_CNP}}": "2900101123456", "{{ADMINISTRATOR_1_NUME}}": "Popescu", "{{ADMINISTRATOR_1_PRENUME}}": "Ion",
        "{{CAMP_INCEPAND_CU}}": "10.11.2025",
    }
    manual = next(f["tag"] for f in found if f["scope"] == "manual")
    values[manual] = "10.11.2025"
    out = fill_docx(template, values)
    text = "\n".join(p.text for p in Document(io.BytesIO(out)).paragraphs)
    assert "{{" not in text and "……" not in text.replace("Semnătura", "")
    assert "domiciliată în Cluj, str. Test 2, jud. Cluj" in text
    assert "seria CJ nr. 123456 eliberată de SPCLEP Cluj la 01.02.2020, CNP 2900101123456" in text
    assert "cu sediul in Timișoara, Str. Exemplu nr. 1, jud. Timiș, C.U.I.:12345678" in text
    assert "reprezentată de administrator Popescu Ion" in text and "începând" not in text and "incepand cu 10.11.2025" in text
    assert "Capitalul social subscris este în valoare de 200 lei, împărțit în 20 părți sociale." in text


def test_rutele_verifica_fisierul_si_alegerile(monkeypatch):
    import app as app_module
    monkeypatch.setattr(app_module, "_verify", lambda *a, **k: "uid")
    c = app_module.app.test_client()
    ok = c.post("/template/blanks", data={"template": (io.BytesIO(make_doc()), "t.docx")}, content_type="multipart/form-data")
    assert ok.status_code == 200 and len(ok.get_json()["blanks"]) >= 15 and "SOCIETATE_SEDIU" in ok.get_json()["companyFields"]
    bad_ext = c.post("/template/blanks", data={"template": (io.BytesIO(b"x"), "t.doc")}, content_type="multipart/form-data")
    assert bad_ext.status_code == 400
    assert c.post("/template/blanks").status_code == 400
    ids = {str(b["id"]): b["tag"] for b in ok.get_json()["blanks"][:3]}
    applied = c.post("/template/blanks/apply", data={"template": (io.BytesIO(make_doc()), "t.docx"), "choices": json.dumps({**ids, "0": "<script>"})},
                     content_type="multipart/form-data")
    assert applied.status_code == 200
    assert "<script>" not in "\n".join(p.text for p in Document(io.BytesIO(applied.data)).paragraphs)   # etichetele nevalide sunt ignorate
    broken = c.post("/template/blanks/apply", data={"template": (io.BytesIO(make_doc()), "t.docx"), "choices": "nu e json"}, content_type="multipart/form-data")
    assert broken.status_code == 400


LOCAL = Path(__file__).resolve().parent.parent / "fisiere_template" / "SabloaneAvocatTavi" / "anonimizate"


@pytest.mark.skipif(not (LOCAL / "ACTUL CONSTITUTIV.docx").exists(), reason="documentele locale ale avocatului nu sunt în repository")
def test_actul_constitutiv_real_e_recunoscut_aproape_complet():
    found = blanks.analyze((LOCAL / "ACTUL CONSTITUTIV.docx").read_bytes())
    manual = [f for f in found if f["scope"] == "manual"]
    assert len(found) > 60 and len(manual) <= 3, (len(found), [m["tag"] for m in manual])
