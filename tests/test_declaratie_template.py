"""Declarația pe propria răspundere: în documentul generat nu trebuie să rămână etichete brute ({{...}}).

Frontend-ul trimite „-” pentru câmpurile necompletate (vezi withEmptyFieldMarks); serverul exclude valorile goale
din înlocuire, deci o valoare goală ar lăsa eticheta în text. Testul folosește șablonul real din registru.
"""
import io
import re
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from doc_filler import fill_docx, list_placeholders_in_docx  # noqa: E402

TEMPLATE = Path(__file__).resolve().parent.parent / "fisiere_template" / "Declarație proprie răspundere desfășurare activități.docx"
ROW_TAGS = {"{{ADRESA}}", "{{CAEN}}", "{{CAEN_DESC}}", "{{NR_CRT}}"}


def _all_text(doc_bytes: bytes) -> str:
    from doc_filler import _iter_all_tables
    doc = Document(io.BytesIO(doc_bytes))
    parts = [p.text for p in doc.paragraphs]
    for t in _iter_all_tables(doc):
        for row in t.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def _flat_replacements(values: dict[str, str]) -> dict[str, str]:
    """Ca serverul: cheile cu valoare goală nu ajung în înlocuire."""
    return {k: v for k, v in values.items() if v}


def _payload(fill_empty: bool) -> tuple[dict[str, str], dict]:
    tags = [t for t in list_placeholders_in_docx(TEMPLATE.read_bytes()) if t not in ROW_TAGS]
    values = {t: "Valoare" for t in tags}
    for t in ("{{DECLARANT_SC}}", "{{DECLARANT_ET}}", "{{SEDIU_ET}}", "{{SEDIU_TEL}}", "{{SOCIETATE_NR_REG}}"):
        values[t] = "-" if fill_empty else ""
    rows = {
        "CAEN_SEDIU": [{"CAEN": "4711", "CAEN_DESC": "Comerț"}],
        "CAEN_TERTI": [],
        "SEDII_SECUNDARE": [
            {"NR_CRT": "1", "ADRESA": "Str. Exemplu 1", "CAEN": "4711", "CAEN_DESC": "Comerț"},
            {"NR_CRT": "", "ADRESA": "", "CAEN": "4719", "CAEN_DESC": "Alt comerț"},
        ],
    }
    return values, rows


def test_empty_values_left_as_raw_tags_without_the_dash(monkeypatch):
    """Documentează problema: fără „-”, câmpurile goale rămân ca etichete în document."""
    values, rows = _payload(fill_empty=False)
    out = fill_docx(TEMPLATE.read_bytes(), _flat_replacements(values), None, None, rows)
    assert "{{DECLARANT_SC}}" in _all_text(out)


def test_dash_leaves_no_raw_tags_and_keeps_blank_row_cells():
    values, rows = _payload(fill_empty=True)
    out = fill_docx(TEMPLATE.read_bytes(), _flat_replacements(values), None, None, rows)
    text = _all_text(out)
    assert not re.search(r"\{\{[^}]*\}\}", text), re.findall(r"\{\{[^}]*\}\}", text)
    assert "Str. Exemplu 1" in text and "4719" in text        # rândurile tabelului s-au completat


def test_act_constitutiv_durata_societatii_fara_data_inmatricularii():
    """Art. „Durata societății”: 99 de ani, fără „începând cu data înmatriculării” (formularea cerută de proprietar)."""
    text = _all_text((TEMPLATE.parent / "Act Constitutiv PJ - actualizare.docx").read_bytes())
    assert "Durata de funcţionare a societăţii este pe o perioadă de 99 de ani." in text
    assert "perioada de 99 de ani, începând cu data înmatriculării" not in text


def test_clauzele_de_durata_fara_perioada_nelimitata():
    """Clauzele de modificare a duratei (societate, mandat administrator) spun doar „într-o perioadă de 99 de ani”."""
    base = TEMPLATE.parent
    for name in ("2 Decizia Asociatului unic.docx", "3 Hot AGA.docx"):
        text = _all_text((base / name).read_bytes())
        assert "Se aprobă modificarea duratei societății într-o perioadă de 99 de ani." in text, name
        assert "Se aprobă modificarea duratei mandatului de administrator într-o perioadă de 99 de ani." in text, name
        assert "dintr-o perioadă nelimitată" not in text, name


def test_clauza_actualizare_caen_rev3_cu_data_de_azi_si_liste_de_coduri():
    """„Actualizare cod CAEN REV3”: data de azi, „astfel:”, apoi codul principal și codurile secundare pe câte un rând."""
    base = TEMPLATE.parent
    for name in ("2 Decizia Asociatului unic.docx", "3 Hot AGA.docx"):
        text = _all_text((base / name).read_bytes())
        assert "din data de {{DATA_AZI}}, astfel:" in text, name
        assert "Activitatea principală:\t{{CAEN_PRINCIPAL_COD}}" in text, name
        assert "Activitatea secundară:\t{{CAEN_SECUNDARE_COD}}" in text, name
        assert "conform REV 3, potrivit Actului Constitutiv din data de {{DATA_CURENTA}}" not in text, name


def test_clauza_caen_rev3_se_completeaza_la_generare():
    import re
    from doc_filler import list_clauses_in_docx
    base = TEMPLATE.parent
    for name in ("2 Decizia Asociatului unic.docx", "3 Hot AGA.docx"):
        raw = (base / name).read_bytes()
        rev3 = next(c for c in list_clauses_in_docx(raw) if "REV3" in c["label"].upper().replace(" ", ""))
        assert set(rev3["placeholders"]) == {"{{DATA_AZI}}", "{{CAEN_PRINCIPAL_COD}}", "{{CAEN_SECUNDARE_COD}}"}, name
        out = fill_docx(raw, {"{{DATA_AZI}}": "15.07.2026", "{{CAEN_PRINCIPAL_COD}}": "7020", "{{CAEN_SECUNDARE_COD}}": "4100, 4311, 4312, 7311"},
                        None, [rev3["tag"]], None)
        text = _all_text(out)
        assert "din data de 15.07.2026, astfel:" in text, name
        assert "Activitatea principală:\t7020" in text and "Activitatea secundară:\t4100, 4311, 4312, 7311" in text, name
        assert not re.search(r"\{\{(DATA_AZI|CAEN_PRINCIPAL_COD|CAEN_SECUNDARE_COD)\}\}", text), name


def test_liniile_cu_coduri_caen_nu_sunt_bold():
    """„Activitatea principală/secundară” și codurile CAEN se scriu normal (doar „Art. N” rămâne bold, ca la celelalte clauze)."""
    base = TEMPLATE.parent
    for name in ("2 Decizia Asociatului unic.docx", "3 Hot AGA.docx"):
        doc = Document(io.BytesIO((base / name).read_bytes()))
        lines = [p for p in doc.paragraphs if p.text.lstrip("\t").startswith(("Activitatea principală:", "Activitatea secundară:"))]
        assert len(lines) == 2, name
        for p in lines:
            assert all(not r.bold for r in p.runs), (name, p.text)


def test_liniile_cu_coduri_caen_sunt_aliniate_cu_indentare_suspendata():
    """Eticheta începe la un tab de margine, iar valorile (și rândurile următoare, dacă lista e lungă) încep aliniat, după etichetă."""
    base = TEMPLATE.parent
    for name in ("2 Decizia Asociatului unic.docx", "3 Hot AGA.docx"):
        doc = Document(io.BytesIO((base / name).read_bytes()))
        lines = [p for p in doc.paragraphs if p.text.startswith(("Activitatea principală:", "Activitatea secundară:"))]
        assert len(lines) == 2, name
        for p in lines:
            pf = p.paragraph_format
            assert pf.left_indent.cm > pf.left_indent.cm + pf.first_line_indent.cm > 0, name     # indentare suspendată
            assert abs(pf.left_indent.cm + pf.first_line_indent.cm - 1.27) < 0.05, name          # eticheta începe la un tab
            assert any(abs(t.position.cm - pf.left_indent.cm) < 0.02 for t in pf.tab_stops), name  # tab-ul duce la valori
        assert lines[0].paragraph_format.left_indent == lines[1].paragraph_format.left_indent, name  # valorile sunt aliniate între ele


def test_tabul_dintre_eticheta_si_coduri_ramane_dupa_completare():
    from doc_filler import list_clauses_in_docx
    raw = (TEMPLATE.parent / "2 Decizia Asociatului unic.docx").read_bytes()
    rev3 = next(c for c in list_clauses_in_docx(raw) if "REV3" in c["label"].upper().replace(" ", ""))
    out = fill_docx(raw, {"{{DATA_AZI}}": "15.07.2026", "{{CAEN_PRINCIPAL_COD}}": "7020", "{{CAEN_SECUNDARE_COD}}": "4100, 4311"}, None, [rev3["tag"]], None)
    text = _all_text(out)
    assert "Activitatea principală:\t7020" in text and "Activitatea secundară:\t4100, 4311" in text


def test_health_arata_versiunile_sabloanelor_de_baza():
    import json
    import app as app_module
    body = app_module.app.test_client().get("/health").get_json()
    reg = json.loads((TEMPLATE.parent / "registry.json").read_text(encoding="utf-8"))
    reg = reg if isinstance(reg, list) else reg.get("templates", reg)
    assert body["status"] == "ok"
    assert body["templates"] == {t["key"]: t["version"] for t in reg}
