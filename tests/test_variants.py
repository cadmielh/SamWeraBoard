"""Variantele „a/b” din documente (sex, număr, categorie) se aleg automat; ce nu se poate stabili rămâne neschimbat."""
import io
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import variants  # noqa: E402
from doc_filler import fill_docx  # noqa: E402

TPL = Path(__file__).resolve().parent.parent / "fisiere_template"


def resolve(text: str, ctx: dict, repl: dict | None = None, fixed=None, split: list[str] | None = None):
    """Rulează motorul pe un paragraf (opțional împărțit în run-uri) și întoarce (text, persoane cu sex necunoscut)."""
    doc = Document()
    p = doc.add_paragraph()
    for part in (split or [text]):
        p.add_run(part)
    assert "".join(split or [text]) == text
    ch = variants.resolve_paragraph(p, ctx, repl or {}, fixed)
    return p.text, variants.unresolved_persons(ch)


M = {"sex": {"ASOCIAT_1": "M"}}
F = {"sex": {"ASOCIAT_1": "F"}}


# ── sex ──────────────────────────────────────────────────────────────────────
def test_sex_dupa_persoana_din_paragraf():
    t = "Subsemnata/ul {{ASOCIAT_1_NUME}}, născut/ă la data de 01.01.1990, domiciliat/ă în Timișoara."
    assert resolve(t, M)[0] == "Subsemnatul {{ASOCIAT_1_NUME}}, născut la data de 01.01.1990, domiciliat în Timișoara."
    assert resolve(t, F)[0] == "Subsemnata {{ASOCIAT_1_NUME}}, născută la data de 01.01.1990, domiciliată în Timișoara."


def test_doua_persoane_in_acelasi_paragraf_fiecare_cu_forma_ei():
    t = ("Domnul/Doamna {{ADMINISTRATOR_VECHI_NUME}} se retrage din calitatea de administrator. "
         "Domnul/Doamna {{ADMINISTRATOR_NOU_NUME}} este numit/ă administrator.")
    ctx = {"sex": {"ADMINISTRATOR_VECHI": "M", "ADMINISTRATOR_NOU": "F"}}
    out, unknown = resolve(t, ctx)
    assert out == ("Domnul {{ADMINISTRATOR_VECHI_NUME}} se retrage din calitatea de administrator. "
                   "Doamna {{ADMINISTRATOR_NOU_NUME}} este numită administrator.")
    assert not unknown


def test_cuvinte_intregi_nascut_nascuta_in_orice_ordine():
    tx = "{{ASOCIAT_1_NUME}}, născut/născută la X, domiciliată/domiciliat în Y, identificat/identificată, Subsemnatul/Subsemnata"
    assert resolve(tx, M)[0] == "{{ASOCIAT_1_NUME}}, născut la X, domiciliat în Y, identificat, Subsemnatul"
    assert resolve(tx, F)[0] == "{{ASOCIAT_1_NUME}}, născută la X, domiciliată în Y, identificată, Subsemnata"
    out, unknown = resolve(tx, {"sex": {"ASOCIAT_1": None}})
    assert out == tx and unknown == {"ASOCIAT_1"}


def test_cuvinte_intregi_nu_confunda_alte_perechi():
    tx = "{{ASOCIAT_1_NUME}} și/sau soțul/soția, emis/emisă, social/profesional, unul/una"
    assert resolve(tx, F)[0] == "{{ASOCIAT_1_NUME}} și/sau soțul/soția, emis/emisă, social/profesional, unul/una"


def test_forme_genitiv_si_ordine_inversa():
    ctx = {"sex": {"ASOCIAT_1": "F"}}
    assert resolve("hotărârea doamnei/domnului {{ASOCIAT_1_NUME}}", ctx)[0] == "hotărârea doamnei {{ASOCIAT_1_NUME}}"
    assert resolve("hotărârea domnului/doamnei {{ASOCIAT_1_NUME}}", {"sex": {"ASOCIAT_1": "M"}})[0] == "hotărârea domnului {{ASOCIAT_1_NUME}}"


def test_forme_cu_paranteza_din_declaratie_dar_nu_cele_despre_document():
    t = "Subsemnatul(a), {{DECLARANT_NUME}}, domiciliat(ă) în {{DECLARANT_LOCALITATE}}, născut(ă) în X, emis(ă) de {{DECLARANT_ACT_EMIS_DE}}"
    out, _ = resolve(t, {"sex": {"DECLARANT": "F"}})
    assert out == "Subsemnata, {{DECLARANT_NUME}}, domiciliată în {{DECLARANT_LOCALITATE}}, născută în X, emis(ă) de {{DECLARANT_ACT_EMIS_DE}}"


def test_sex_necunoscut_nu_se_ghiceste_si_e_raportat():
    t = "{{ASOCIAT_1_NUME}}, născut/ă la data de 01.01.1990"
    out, unknown = resolve(t, {"sex": {"ASOCIAT_1": None}})
    assert out == t and unknown == {"ASOCIAT_1"}


def test_majuscule_la_inceput_de_propozitie():
    assert resolve("Domnul/Doamna {{ASOCIAT_1_NUME}}", F)[0] == "Doamna {{ASOCIAT_1_NUME}}"
    assert resolve("NUMIT/Ă {{ASOCIAT_1_NUME}}", F)[0] == "NUMITĂ {{ASOCIAT_1_NUME}}"


def test_alternativa_impartita_pe_mai_multe_run_uri_pastreaza_formatarea():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("{{ASOCIAT_1_NUME}} este ")
    p.add_run("numit").bold = True
    p.add_run("/ă administrator")
    variants.resolve_paragraph(p, F, {})
    assert p.text == "{{ASOCIAT_1_NUME}} este numită administrator"
    assert p.runs[1].bold is True


# ── număr ────────────────────────────────────────────────────────────────────
def test_asociat_unic_sau_asociati():
    t = "având ca asociat unic/asociați pe: hotărârea asociatului\nunic/asociaților. Asociatul/ asociații nemulțumit/nemulțumiți poate/pot face opoziție."
    one, _ = resolve(t, {"asociati": 1})
    assert one == "având ca asociat unic pe: hotărârea asociatului\nunic. Asociatul nemulțumit poate face opoziție."
    many, _ = resolve(t, {"asociati": 3})
    assert many == "având ca asociați pe: hotărârea asociaților. Asociații nemulțumiți pot face opoziție."


def test_administrator_sau_administratori_si_verbe():
    t = "Administratorul/Administratorii reprezintă societatea. Asociatul unic/ asociații îşi va/vor putea angaja un administrator."
    assert resolve(t, {"asociati": 1, "administratori": 1})[0] == "Administratorul reprezintă societatea. Asociatul unic îşi va putea angaja un administrator."
    assert resolve(t, {"asociati": 2, "administratori": 2})[0] == "Administratorii reprezintă societatea. Asociații îşi vor putea angaja un administrator."


def test_majuscule_complete_la_titluri():
    assert resolve("ASOCIAT UNIC/ASOCIAȚI:", {"asociati": 1})[0] == "ASOCIAT UNIC:"
    assert resolve("ASOCIAT UNIC/ASOCIAȚI:", {"asociati": 2})[0] == "ASOCIAȚI:"


def test_numar_necunoscut_lasa_alternativa():
    t = "asociat unic/asociați"
    assert resolve(t, {})[0] == t
    assert resolve(t, {"asociati": 0})[0] == t


def test_perechi_de_numar_cu_pluralul_primul():
    t = "În calitate de asociat unic/asociați {{ASOCIATI_LISTA}} au/are drepturile."
    assert resolve(t, {"asociati": 1})[0] == "În calitate de asociat unic {{ASOCIATI_LISTA}} are drepturile."
    assert resolve(t, {"asociati": 3})[0] == "În calitate de asociați {{ASOCIATI_LISTA}} au drepturile."


# ── categorie ────────────────────────────────────────────────────────────────
def test_social_sau_profesional_dupa_tipul_clientului():
    t = "sediul social/profesional"
    assert resolve(t, {"tip": "PJ"})[0] == "sediul social"
    assert resolve(t, {"tip": "PF"})[0] == "sediul profesional"
    assert resolve(t, {})[0] == t
    assert resolve("SOCIAL/PROFESIONAL", {"tip": "PF"})[0] == "PROFESIONAL"


def test_social_profesional_secundar_ramane():
    t = "în afara sediului social/profesional/secundar"
    assert resolve(t, {"tip": "PJ"})[0] == t


def test_judet_sau_sector_dupa_valoarea_din_document():
    t = "județ/sector {{SEDIU_JUDET}}, sector/județ {{DECLARANT_NASTERE_JUDET}}"
    out, _ = resolve(t, {}, {"{{SEDIU_JUDET}}": "București", "{{DECLARANT_NASTERE_JUDET}}": "Timiș"})
    assert out == "sector {{SEDIU_JUDET}}, județ {{DECLARANT_NASTERE_JUDET}}"
    assert resolve(t, {}, {"{{SEDIU_JUDET}}": "-"})[0].startswith("județ/sector")   # necompletat → rămâne


def test_cnp_sau_nif():
    t = "CNP/NIF {{DECLARANT_CNP}}"
    assert resolve(t, {}, {"{{DECLARANT_CNP}}": "1900101123456"})[0] == "CNP {{DECLARANT_CNP}}"
    assert resolve(t, {}, {"{{DECLARANT_CNP}}": "9900101123456"})[0] == "NIF {{DECLARANT_CNP}}"
    assert resolve(t, {}, {})[0] == t


def test_alte_alternative_nu_se_ating():
    t = "și/sau funcționare/desfășurare județ/ora"
    assert resolve(t, {"tip": "PJ", "asociati": 1})[0] == t


# ── în documente reale ───────────────────────────────────────────────────────
def _text(doc_bytes: bytes) -> str:
    from doc_filler import _iter_all_tables
    d = Document(io.BytesIO(doc_bytes))
    parts = [p.text for p in d.paragraphs]
    for t in _iter_all_tables(d):
        for r in t.rows:
            for c in r.cells:
                parts.extend(p.text for p in c.paragraphs)
    return "\n".join(parts)


def test_bloc_repetitiv_fiecare_asociat_cu_sexul_lui():
    doc = Document()
    for t in ("{{#ASOCIATI}}", "{{INDEX}}. {{NUME}}, născut/ă la X, domiciliat/ă în Y", "{{/ASOCIATI}}"):
        doc.add_paragraph(t)
    buf = io.BytesIO(); doc.save(buf)
    groups = {"ASOCIATI": [
        {"NUME": "Popescu", "SEX": "M", "SEX_REF": "ASOCIAT_1"},
        {"NUME": "Ionescu", "SEX": "F", "SEX_REF": "ASOCIAT_2"},
        {"NUME": "Necunoscut", "SEX": "", "SEX_REF": "ASOCIAT_3"},
    ]}
    warnings: set = set()
    out = _text(fill_docx(buf.getvalue(), {}, groups, None, None, {"sex": {}}, warnings))
    assert "1. Popescu, născut la X, domiciliat în Y" in out
    assert "2. Ionescu, născută la X, domiciliată în Y" in out
    assert "3. Necunoscut, născut/ă la X, domiciliat/ă în Y" in out
    assert warnings == {"ASOCIAT_3"}


def test_actul_constitutiv_real_asociat_unic_vs_asociati():
    raw = (TPL / "Act Constitutiv PJ - actualizare.docx").read_bytes()
    unic = _text(fill_docx(raw, {}, None, None, None, {"asociati": 1, "administratori": 1, "tip": "PJ", "sex": {}}))
    assert "unic/asociați" not in unic and "Administratorul/Administratorii" not in unic and "va/vor" not in unic
    assert "asociat unic pe:" in unic and "Administratorul reprezintă" in unic
    multi = _text(fill_docx(raw, {}, None, None, None, {"asociati": 3, "administratori": 2, "tip": "PJ", "sex": {}}))
    assert "asociați pe:" in multi and "Administratorii reprezintă" in multi and "unic/asociați" not in multi


def test_declaratia_reala_social_profesional_si_forme_de_sex():
    raw = (TPL / "Declarație proprie răspundere desfășurare activități.docx").read_bytes()
    repl = {"{{DECLARANT_CNP}}": "2900101123456", "{{SEDIU_JUDET}}": "București", "{{DECLARANT_JUDET}}": "Timiș",
            "{{DECLARANT_NASTERE_JUDET}}": "Cluj", "{{DECLARANT_NUME}}": "Ionescu"}
    out = " ".join(_text(fill_docx(raw, repl, None, None, {"CAEN_SEDIU": [], "CAEN_TERTI": [], "SEDII_SECUNDARE": []},
                                   {"tip": "PJ", "sex": {"DECLARANT": "F"}})).split())
    assert "Subsemnata, Nume Ionescu" in out and "domiciliată în localitatea" in out and "născută în localitatea" in out
    assert "emis(ă) de" in out                                    # se referă la actul de identitate, nu la persoană
    assert "CNP 2900101123456" in out and "CNP/NIF" not in out    # cetățean român: CNP
    assert "sediul social în:" in out and "sediul social/profesional în" not in out
    assert "sector București" in out                              # sediu în București → „sector”
    assert "județ Timiș" in out and "județ Cluj" in out           # în rest → „județ”


def test_ctx_din_client_este_validat():
    ctx = variants.sanitize_ctx({"sex": {"ASOCIAT_1": "x", "bad key!": "M", "OK": "F"}, "asociati": "3", "administratori": 2, "tip": "??"})
    assert ctx == {"sex": {"ASOCIAT_1": None, "OK": "F"}, "asociati": None, "administratori": 2, "tip": None}
    assert variants.sanitize_ctx("nu e dict") == {}
