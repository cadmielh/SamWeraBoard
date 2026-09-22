"""Șabloane cu poziții numerotate (ASOCIAT_1_..., ASOCIAT_2_...) scrise pentru un număr fix de persoane:
niciodată nu trebuie să dispară în tăcere o persoană — nici cea deja rezolvată dintr-un paragraf mixt, nici
cea care depășește numărul de poziții scrise în șablon."""
import io
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from doc_filler import fill_docx  # noqa: E402


def _doc(text: str) -> bytes:
    d = Document()
    d.add_paragraph(text)
    b = io.BytesIO(); d.save(b)
    return b.getvalue()


MIXED = "Constituit de asociatii {{ASOCIAT_1_NUME}} {{ASOCIAT_1_PRENUME}}, CNP {{ASOCIAT_1_CNP}} si {{ASOCIAT_2_NUME}} {{ASOCIAT_2_PRENUME}}, CNP {{ASOCIAT_2_CNP}}."


def test_un_singur_asociat_real_pastreaza_datele_lui_si_avertizeaza():
    """Șablonul scris pentru 2: cu un singur asociat real, paragraful NU se șterge (ar pierde și datele valide)."""
    w = set()
    out = fill_docx(_doc(MIXED), {"{{ASOCIAT_1_NUME}}": "Popescu", "{{ASOCIAT_1_PRENUME}}": "Ion", "{{ASOCIAT_1_CNP}}": "1900101123456"},
                    capacity_warnings=w)
    text = Document(io.BytesIO(out)).paragraphs[0].text
    assert "Popescu Ion" in text and "1900101123456" in text     # datele reale n-au dispărut
    assert "{{ASOCIAT_2_NUME}}" in text                          # poziția lipsă rămâne vizibilă, nu ascunsă
    assert len(w) == 1 and "1 asociați" in next(iter(w))


def test_trei_asociati_reali_template_scris_pentru_doi_avertizeaza_pe_al_treilea():
    w = set()
    out = fill_docx(_doc(MIXED), {
        "{{ASOCIAT_1_NUME}}": "Popescu", "{{ASOCIAT_1_PRENUME}}": "Ion", "{{ASOCIAT_1_CNP}}": "1900101123456",
        "{{ASOCIAT_2_NUME}}": "Ionescu", "{{ASOCIAT_2_PRENUME}}": "Maria", "{{ASOCIAT_2_CNP}}": "2900101123456",
        "{{ASOCIAT_3_NUME}}": "Vasile", "{{ASOCIAT_3_PRENUME}}": "Dan", "{{ASOCIAT_3_CNP}}": "1900101123457",
    }, capacity_warnings=w)
    text = Document(io.BytesIO(out)).paragraphs[0].text
    assert "Popescu Ion" in text and "Ionescu Maria" in text
    assert len(w) == 1 and "sunt 3" in next(iter(w)) and "2" in next(iter(w))


def test_numarul_exact_potrivit_nu_produce_niciun_avertisment():
    w = set()
    out = fill_docx(_doc(MIXED), {
        "{{ASOCIAT_1_NUME}}": "Popescu", "{{ASOCIAT_1_PRENUME}}": "Ion", "{{ASOCIAT_1_CNP}}": "1900101123456",
        "{{ASOCIAT_2_NUME}}": "Ionescu", "{{ASOCIAT_2_PRENUME}}": "Maria", "{{ASOCIAT_2_CNP}}": "2900101123456",
    }, capacity_warnings=w)
    text = Document(io.BytesIO(out)).paragraphs[0].text
    assert "{{" not in text and w == set()


def test_marcajul_intern_nu_ajunge_in_fisierul_salvat():
    import zipfile
    out = fill_docx(_doc(MIXED), {"{{ASOCIAT_1_NUME}}": "Popescu", "{{ASOCIAT_1_PRENUME}}": "Ion", "{{ASOCIAT_1_CNP}}": "1900101123456"},
                    capacity_warnings=set())
    xml = zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml").decode()
    assert "swbProtected" not in xml


def test_sablon_cu_bloc_repetitiv_nu_declanseaza_avertismente_false():
    """Un șablon care folosește deja {{#ASOCIATI}} (scalabil) nu trebuie să pară „lipsă de poziții”."""
    d = Document()
    d.add_paragraph("{{#ASOCIATI}}")
    d.add_paragraph("{{INDEX}}. {{NUME}} {{PRENUME}}")
    d.add_paragraph("{{/ASOCIATI}}")
    b = io.BytesIO(); d.save(b)
    w = set()
    groups = {"ASOCIATI": [{"NUME": "A", "PRENUME": "a"}, {"NUME": "B", "PRENUME": "b"}, {"NUME": "C", "PRENUME": "c"}]}
    out = fill_docx(b.getvalue(), {"{{ASOCIAT_1_NUME}}": "A", "{{ASOCIAT_2_NUME}}": "B", "{{ASOCIAT_3_NUME}}": "C"},
                    groups=groups, capacity_warnings=w)
    text = "\n".join(p.text for p in Document(io.BytesIO(out)).paragraphs)
    assert "A a" in text and "B b" in text and "C c" in text
    assert w == set()


def test_fara_replacements_numerotate_nu_calculeaza_nimic():
    w = set()
    out = fill_docx(_doc("Text simplu {{CNP}}"), {"{{CNP}}": "1900101123456"}, capacity_warnings=w)
    assert w == set() and "1900101123456" in Document(io.BytesIO(out)).paragraphs[0].text


def test_capacity_warnings_optional_nu_e_obligatoriu():
    # apel fără parametrul nou — comportament neschimbat pentru orice cod vechi care nu-l trimite
    out = fill_docx(_doc(MIXED), {"{{ASOCIAT_1_NUME}}": "Popescu", "{{ASOCIAT_1_PRENUME}}": "Ion", "{{ASOCIAT_1_CNP}}": "1900101123456"})
    text = Document(io.BytesIO(out)).paragraphs[0].text
    assert "Popescu Ion" in text  # tot nu se pierde, chiar fără să colectăm avertismentul


def test_ruta_fill_docx_expune_avertismentele_de_capacitate(monkeypatch):
    import app as app_module
    monkeypatch.setattr(app_module, "_verify", lambda *a, **k: "uid")
    c = app_module.app.test_client()
    r = c.post("/fill/docx", data={
        "template": (io.BytesIO(_doc(MIXED)), "t.docx"),
        "{{ASOCIAT_1_NUME}}": "Popescu", "{{ASOCIAT_1_PRENUME}}": "Ion", "{{ASOCIAT_1_CNP}}": "1900101123456",
    }, content_type="multipart/form-data")
    assert r.status_code == 200
    import json as _json
    warnings = _json.loads(r.headers["X-Capacity-Warnings"])
    assert len(warnings) == 1 and "1 asociați" in warnings[0]
