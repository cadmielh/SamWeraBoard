"""Șabloane fără etichete: locurile libere („……”, „.....”, „_____”) se recunosc din context și se leagă de câmpurile aplicației.

Fluxul (vezi rutele /template/blanks și /template/blanks/apply din app.py):
  1. `analyze(docx)` găsește fiecare loc liber, cu textul din jur, și propune câmpul potrivit (ex. „CNP ……” → CNP-ul persoanei,
     „sediul în ……” → sediul societății, „aport de …… lei” → aportul asociatului). Ce nu se recunoaște devine câmp manual.
  2. Utilizatorul confirmă sau corectează propunerile.
  3. `apply(docx, alegeri)` înlocuiește fiecare loc liber cu eticheta aleasă → un șablon obișnuit, pe care motorul existent
     (variante după sex/număr, „-” la câmpuri goale, procent de completare) îl completează.

Regulile sunt pentru formulările uzuale din actele societăților; propunerile sunt doar propuneri: nu se aplică nimic fără confirmare.
"""
from __future__ import annotations

import copy
import io
import itertools
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from variants import _norm, _replace_span

BLANK_RE = re.compile(r"…+|\.{4,}|_{3,}")
_ARTIFACT_DOTS_RE = re.compile(r"\.{1,3}")


def _extend_past_artifact_dots(text: str, end: int) -> int:
    """„….,” / „….. jud” — puncte literale rămase lipite de o elipsă („…”): BLANK_RE prinde doar elipsa în
    sine, nu și punctele care mai rămân când autorul a tastat mai multe puncte decât a convertit Word automat
    (de obicei primele trei) în elipsă. Dacă imediat după aceste puncte urmează o literă mică sau o virgulă
    (propoziția chiar continuă), punctele fac parte din locul liber, nu sunt punctuație reală — altfel ar
    rămâne vizibile în documentul completat („cetătean romana., născut”). Dacă urmează literă mare sau
    paragraful se termină acolo, punctele sunt sfârșit de propoziție — rămân neschimbate."""
    m = _ARTIFACT_DOTS_RE.match(text, end)
    if not m:
        return end
    j = m.end()
    while j < len(text) and text[j] == " ":     # un spațiu între puncte și cuvântul următor nu schimbă nimic
        j += 1
    nxt = text[j:j + 1]
    if nxt == "," or (nxt.isalpha() and nxt.islower()):
        return m.end()
    return end

# Câmpuri ale societății (etichetele există deja în aplicație)
COMPANY_FIELDS = {
    "SOCIETATE_DENUMIRE": "Denumirea societății",
    "SOCIETATE_SEDIU": "Sediul societății (adresa completă)",
    "SOCIETATE_SEDIU_FARA_JUDET": "Sediul societății (fără județ — folosit când județul e un loc liber separat)",
    "SOCIETATE_JUDET": "Județul sediului",
    "SOCIETATE_CIF": "CUI / cod fiscal",
    "SOCIETATE_NR_REG": "Nr. de înmatriculare (J…/…/…)",
    "CAPITAL_SOCIAL_TOTAL": "Capitalul social (lei)",
    "PARTI_SOCIALE_TOTALE": "Numărul total de părți sociale",
    "CAEN_1": "Activitatea principală (cod și denumire)",
    "CAEN_DOMENIU": "Domeniul principal de activitate (grupa CAEN, cod și denumire)",
    "CAEN_PRINCIPAL_COD": "Cod CAEN principal",
    "CAEN": "O activitate secundară (cod și denumire)",
    "DATA_AZI": "Data de azi",
    "ASOCIATI_LISTA": "Numele tuturor asociaților, pe o linie („A, B și C”)",
    "ADMINISTRATORI_LISTA": "Numele tuturor administratorilor, pe o linie",
}

# Câmpuri per persoană: eticheta finală = {{ASOCIAT_n_CÂMP}} / {{ADMINISTRATOR_n_CÂMP}}
PERSON_FIELDS = {
    "NUME_COMPLET": "Nume și prenume",
    "NUME": "Nume",
    "PRENUME": "Prenume",
    "CNP": "CNP",
    "ADRESA": "Adresa de domiciliu",
    "JUDET": "Județul",
    "DATA_NASTERII": "Data nașterii",
    "LOCUL_NASTERII": "Locul nașterii",
    "CETATENIA": "Cetățenia",
    "SERIE_ACT": "Seria actului de identitate",
    "NR_ACT": "Numărul actului de identitate",
    "SERIE_NUMAR": "Seria și numărul actului",
    "EMISA_DE": "Actul eliberat de",
    "VALABILA_DE_LA": "Data eliberării actului",
    "VALABILA_PANA_LA": "Valabil până la",
    "COTA_PARTICIPARE": "Cota de participare",
    "CAPITAL_SOCIAL": "Aportul la capital (lei)",
    "PARTI_SOCIALE": "Numărul de părți sociale",
}

_MANUAL_PREFIX = "CAMP_"


def person_tag(role: str, n: int, field: str) -> str:
    """Eticheta finală pentru un câmp de persoană; câmpurile de capital folosesc formele existente (CAPITAL_SOCIAL_ASOCIAT_n)."""
    if field == "NUME_COMPLET":
        return f"{{{{{role}_{n}_NUME}}}} {{{{{role}_{n}_PRENUME}}}}"
    if field == "CAPITAL_SOCIAL":
        return f"{{{{CAPITAL_SOCIAL_ASOCIAT_{n}}}}}"
    if field == "PARTI_SOCIALE":
        return f"{{{{PARTI_SOCIALE_ASOCIAT_{n}}}}}"
    return f"{{{{{role}_{n}_{field}}}}}"


def manual_tag(label: str) -> str:
    words = re.sub(r"[^A-Z0-9]+", "_", _norm(label).upper()).strip("_")
    return "{{" + _MANUAL_PREFIX + (words[:40] or "VALOARE") + "}}"


def _paragraphs(doc):
    """Aceeași ordine la analiză și la aplicare: corp, apoi tabele (inclusiv imbricate)."""
    out = list(doc.paragraphs)
    seen: set = set()
    stack = list(doc.tables)
    while stack:
        t = stack.pop(0)
        for r in t.rows:
            for c in r.cells:
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                out.extend(c.paragraphs)
                stack.extend(c.tables)
    return out


class _State:
    """Ce s-a decis mai devreme în același paragraf (persoana curentă și câmpul anterior), pentru legături între locuri libere."""
    def __init__(self):
        self.person = 0          # persoanele „definite” (nume urmat de date de identitate)
        self.mention = 0         # persoanele doar amintite (liste de nume: „asociații X și Y”)
        self.last_field: str | None = None
        self.last_scope: str | None = None
        # Numărul persoanei „curente” — fie definită complet, fie doar amintită — la care se leagă orice alt
        # câmp de persoană care urmează imediat (CNP, aport, cotă…), indiferent care din cele două contoare
        # de mai sus a produs numărul. Fără asta, un câmp legat de o persoană doar AMINTITĂ (ex. „Asociatul
        # …… contribuie cu un aport de …… lei”) ar folosi din greșeală `self.person` (nemodificat de mențiuni),
        # nu numărul mențiunii — aportul celei de-a doua persoane amintite ar ieși etichetat tot ca a primei.
        self.last_person_n = 0


# Calitatea juridică a unei persoane, recunoscută din context — nu doar Asociat/Administrator, ca detectarea
# să funcționeze generic pe orice tip de act (contract de comodat, împuternicire…), nu doar pe cele de
# societate. Lista de mai jos acoperă calitățile uzuale întâlnite până acum; poate crește oricând cu altele
# noi, fără să afecteze restul motorului — fiecare intrare mapează un cuvânt-cheie (fără sufixul de
# declinare, pe care \w*-ul din regex-urile de mai sus îl înghite la potrivirea numelui) la eticheta lui
# canonică (folosită direct în etichete: {{COMODANT_1_NUME}} etc.).
_ROLE_KEYWORDS: list[tuple[str, str]] = [
    # „administr” (nu „administrator”): prinde și „administrarea/administrării societății vor fi
    # îndeplinite de...” — o formă (substantivul „administrare”, de la verbul „a administra”) care nu conține
    # literal cuvântul „administrator”, dar înseamnă exact același lucru; fără rădăcina scurtă, un paragraf
    # ca acesta cădea pe implicitul „ASOCIAT” — greșit, și cu efect vizibil: la generare, rolul ASOCIAT poate
    # aduna TOȚI asociații firmei (listă scalabilă), nu doar administratorul real.
    ("administr", "ADMINISTRATOR"),
    ("asociat", "ASOCIAT"),
    ("comodant", "COMODANT"),
    ("comodatar", "COMODATAR"),
    ("reprezentant legal", "REPREZENTANT_LEGAL"),
    ("reprezentant", "REPREZENTANT_LEGAL"),
    ("imputernicit", "IMPUTERNICIT"),
    ("mandatar", "MANDATAR"),
    ("chirias", "CHIRIAS"),
    ("locator", "LOCATOR"),
    ("locatar", "LOCATAR"),
    ("vanzator", "VANZATOR"),
    ("cumparator", "CUMPARATOR"),
    ("imprumutator", "IMPRUMUTATOR"),
    ("imprumutat", "IMPRUMUTAT"),
    ("cenzor", "CENZOR"),
    ("actionar", "ACTIONAR"),
    ("fondator", "FONDATOR"),
    ("beneficiar", "BENEFICIAR"),
    ("garant", "GARANT"),
    ("debitor", "DEBITOR"),
    ("creditor", "CREDITOR"),
]

# „Reprezentantul legal al acestei societăți VA FI asociatul ……” — rolul e subiectul propoziției
# (REPREZENTANT_LEGAL), urmat de copula „va fi”/„vor fi”; „asociatul”, deși mai aproape de locul liber, e
# doar cuvântul care descrie CINE ocupă acel rol, nu o declarație de rol în sine. Cea mai lungă potrivire
# (ex. „reprezentant legal” înaintea lui „reprezentant”) e încercată prima — sortare explicită, nu ordinea
# din _ROLE_KEYWORDS (poate diferi). Nu trece de „.”/„;” — o propoziție separată, chiar dacă întâmplător
# conține și un cuvânt de rol și „va fi”, nu se leagă de aceeași declarație.
_ROLE_SUBJECT_RE = re.compile(
    r"\b(" + "|".join(sorted({kw for kw, _ in _ROLE_KEYWORDS}, key=len, reverse=True)) + r")\w*[^.;]{0,60}?\b(?:va|vor)\s+fi\b"
)
_ROLE_BY_KEYWORD = dict(_ROLE_KEYWORDS)


def _last_match(regex: re.Pattern, text: str) -> re.Match | None:
    """Ultima potrivire a lui `regex` în `text` (regex.finditer parcurge deja de la stânga la dreapta — luăm
    ultima, cea mai apropiată de finalul textului), sau None dacă nu se potrivește deloc. Comun pentru
    _role_subject/_explicit_role — amândouă vor cea mai RECENTĂ declarație, nu prima."""
    m = None
    for mm in regex.finditer(text):
        m = mm
    return m


def _role_subject(text: str) -> str | None:
    """Ultimul rol găsit ca subiect + copulă „va/vor fi” în `text` (vezi _ROLE_SUBJECT_RE), sau None."""
    m = _last_match(_ROLE_SUBJECT_RE, _norm(text))
    return _ROLE_BY_KEYWORD[m.group(1)] if m else None

# „…, în calitate de <cuvânt(cuvinte)>…” — tipar generic, are prioritate: prinde ORICE calitate scrisă
# explicit de autorul șablonului, chiar una absentă din lista de mai sus (nu se pot anticipa toate tipurile
# de contract dinainte). Cuvântul prins devine el însuși eticheta rolului (ex. „garant ipotecar” → GARANT_IPOTECAR).
_ROLE_EXPLICIT_RE = re.compile(r"\bin\s+calitate\s+de\s+([a-z][a-z\s]{1,40}?)(?=\s+(?:si|iar|care|ce)\b|\s*[,.:;)]|\s*$)")


def _explicit_role(text: str) -> str | None:
    """Ultima calitate declarată explicit („…, în calitate de <cuvânt(cuvinte)>…”) găsită în `text`, sau None
    — folosit atât înainte de locul liber (vezi _role_for), cât și după (vezi analyze()): actele reale scriu
    des calitatea DUPĂ identitatea completă a persoanei („……, domiciliat…, CNP…, în calitate de COMODANT”),
    nu înainte de nume."""
    m = _last_match(_ROLE_EXPLICIT_RE, _norm(text))
    if not m:
        return None
    word = m.group(1).strip()
    # dacă cuvântul prins e chiar unul dintre cele deja cunoscute (la orice formă — „asociati”/„asociatul”
    # pentru „asociat” etc.), folosim eticheta lui canonică, nu textul literal — altfel „în calitate de
    # asociați” ar deveni un rol NOU, „ASOCIATI”, diferit de „ASOCIAT” deja folosit peste tot în rest.
    for keyword, role in _ROLE_KEYWORDS:
        if word == keyword or word.startswith(keyword + " ") or (word.startswith(keyword) and len(word) - len(keyword) <= 3):
            return role
    return re.sub(r"[^a-z0-9]+", "_", word).strip("_").upper()


def _role_for(before_text: str) -> str:
    """Calitatea persoanei, din tot contextul dinaintea locului liber — vezi _role_for_hint. Fără niciun
    indiciu, rămâne „ASOCIAT” (comportamentul dinainte, implicit pentru actele de societate)."""
    return _role_for_hint(before_text) or "ASOCIAT"


def _all_positions(text: str, sub: str) -> list[int]:
    """Toate pozițiile (nu doar ultima) la care apare `sub` în `text`, inclusiv suprapuse."""
    out = []
    start = 0
    while True:
        idx = text.find(sub, start)
        if idx == -1:
            return out
        out.append(idx)
        start = idx + 1


def _role_for_hint(before_text: str) -> str | None:
    """Calitatea persoanei DOAR dacă există un semnal real în textul dinaintea locului liber — fie explicit
    marcată („în calitate de X”), fie subiect + copulă „va/vor fi” (vezi _role_subject — prioritate maximă,
    alături de „în calitate de X”), fie dintr-un cuvânt cunoscut (vezi _ROLE_KEYWORDS), cel mai APROPIAT de
    locul liber (ultimul găsit) câștigă — un paragraf poate vorbi, pe rând, despre mai multe calități diferite,
    fiecare lângă persoana ei (ex. comodant urmat de comodatar). None dacă nu s-a găsit niciun semnal —
    separată de _role_for (care cade pe implicitul „ASOCIAT” în acest caz), ca analyze() să știe când poate
    suprascrie în siguranță cu o calitate găsită DUPĂ locul liber (vezi clause_role) — un semnal real găsit
    ÎNAINTE (ex. „administrator”) nu trebuie suprascris de o calitate care, mai departe în frază, descrie de
    fapt PARTEA/firma („reprezentată de administrator ……, în calitate de COMODATAR” — COMODATAR e calitatea
    firmei, nu a administratorului care o reprezintă).

    EXCEPȚIE, la fel de generică pentru orice cuvânt de rol, nu doar pentru „asociat”: o apariție prinsă
    într-un marcaj de plural cu bară („asociatul/asociații”, „administratorul/administratorii” — vezi
    _PLURAL_MARKER_RE) e un semnal SLAB — marcajul spune explicit „exemplu scris, poate fi oricare”, nu
    afirmă calitatea reală a persoanei; e adesea folosit doar ca formulă de adresare, chiar și când fraza
    descrie alt rol mai departe („Atribuțiile legate de ADMINISTRAREA societății vor fi îndeplinite de
    asociatul/asociații ……” — cel mai apropiat cuvânt de locul liber e „asociatul”, dintr-un marcaj de plural,
    dar fraza vorbește clar despre administrator). O apariție SLABĂ câștigă doar dacă nu există nicăieri o
    apariție puternică (în afara unui asemenea marcaj) a vreunui cuvânt de rol."""
    explicit = _explicit_role(before_text)
    if explicit:
        return explicit
    subject = _role_subject(before_text)
    if subject:
        return subject
    norm = _norm(before_text)
    weak_spans = [m.span() for m in _PLURAL_MARKER_RE.finditer(norm)]

    def in_weak_span(pos: int, length: int) -> bool:
        return any(s <= pos and pos + length <= e for s, e in weak_spans)

    best_pos, best_role = -1, None
    weak_pos, weak_role = -1, None
    for keyword, role in _ROLE_KEYWORDS:
        for pos in _all_positions(norm, keyword):
            if in_weak_span(pos, len(keyword)):
                if pos > weak_pos:
                    weak_pos, weak_role = pos, role
            elif pos > best_pos:
                best_pos, best_role = pos, role
    return best_role or weak_role


# ── Semnături: linia cu numele persoanei, lângă un titlu/etichetă de secțiune de semnătură
# („SEMNĂTURA ASOCIAT/ASOCIAȚI”, „SEMNATURILE,”, sau doar „Administrator”/„Asociat” singur pe rând) — fie
# tastată ca „……” simplu, singur pe rând (semnătura olografă merge pe rândul gol de dedesubt), fie ca
# „…(indiciu)… _____” (indiciu explicit + linie de subliniere). Fiecare astfel de linie devine o poziție
# numerotată (ASOCIAT_N/ADMINISTRATOR_N, vezi analyze()) — numărată SEPARAT pe rol, ca un asociat urmat mai
# jos de un administrator (persoane diferite) să nu-și „fure” numerele unul altuia.
_ROLE_LABEL_RE = re.compile("|".join(kw.replace(" ", r"\s+") + r"\w*" for kw, _ in _ROLE_KEYWORDS))

# „asociatul ……”, „comodatarul ……”, „reprezentantul legal ……” — un nume e doar AMINTIT (nu introdus complet)
# când în fața lui stă direct un cuvânt de calitate (vezi _ROLE_KEYWORDS) sau un termen generic de adresare
# („subsemnatul”, „domnul”…). Vezi suggest().
_MENTION_TRIGGER_RE = re.compile(rf"\b(?:{_ROLE_LABEL_RE.pattern}|subsemnat\w*|domnul|doamna|dl\.)\s*,?\s*$")


def _signature_role_for_heading(norm_text: str) -> str | None:
    """Rolul unei secțiuni de semnătură, dacă paragraful (deja normalizat: fără diacritice, minuscule, fără
    spații/punctuație la capete) e un titlu care conține „semnatur”, sau doar eticheta unei calități, singură
    pe rând („Administrator”, „Comodatar”…, la orice formă/plural — vezi _ROLE_KEYWORDS) — None altfel."""
    if "semnatur" in norm_text:
        return _role_for(norm_text)
    if _ROLE_LABEL_RE.fullmatch(norm_text):
        return _role_for(norm_text)
    return None


def _looks_like_signature_remainder(s: str) -> bool:
    """Ce rămâne dintr-un paragraf după eliminarea locului liber (+ indiciul din paranteză, dacă exista) —
    dacă e gol sau doar spații/liniuțe de subliniere, paragraful era practic „doar” locul liber, tipic pentru
    o linie de semnătură (numele urmat de o linie de subliniat pentru semnătura olografă)."""
    return not s.strip(" \t_-")


def suggest(before: str, after: str, para_before: str, state: _State, pi: int = 99, alone: bool = False) -> dict:
    """Propunerea pentru un loc liber: {scope: company|person|manual, field, role, person, confidence}.

    `pi` = poziția paragrafului în document, `alone` = locul liber e singur în paragraf (ex. denumirea firmei sub titlu)."""
    b = _norm(before)[-90:]
    a = _norm(after)[:50]
    bt = b.rstrip(" :,;-–")

    def company(field, conf="high"):
        state.last_field, state.last_scope = field, "company"
        return {"scope": "company", "field": field, "confidence": conf}

    def person(field, conf="high", new_person=False, mention=False):
        if field == "NUME_COMPLET":
            if mention:                               # persoană doar amintită: se numără separat de cele definite
                state.mention += 1
                n = state.mention
            else:
                if new_person or state.person == 0:
                    state.person += 1
                n = state.person
            state.last_person_n = n                   # vezi _State.last_person_n
        else:
            # Câmp legat de persoana „curentă” (ultimul NUME_COMPLET din acest paragraf, definit sau doar
            # amintit) — dacă niciun nume n-a apărut încă în paragraf, implicit persoana 1 (comportamentul
            # dinainte, pentru paragrafele care încep direct cu un câmp — CNP, aport… —, nu cu numele).
            n = state.last_person_n or 1
        state.last_field, state.last_scope = field, "person"
        return {"scope": "person", "field": field, "role": _role_for(para_before), "person": n, "confidence": conf}

    if alone and pi <= 4:                              # denumirea firmei, singură pe rând sub titlu
        return company("SOCIETATE_DENUMIRE", "medium")

    # — identitate (câmpuri legate de persoana curentă) —
    if re.search(r"\bcnp\s*:?$", b):
        return person("CNP")
    if re.search(r"\bseria\s*:?$", b):
        return person("SERIE_ACT")
    if re.search(r"\bnr\.?\s*:?$", b) and state.last_field == "SERIE_ACT":
        return person("NR_ACT")
    if re.search(r"eliberat[a]?\s+de\s*$", b):
        return person("EMISA_DE")
    if re.search(r"\bla\s*(?:data\s+de\s*)?$", b) and state.last_field == "EMISA_DE":
        return person("VALABILA_DE_LA")
    if re.search(r"valabil[a]?\s+pana\s+la(?:\s+data\s+de)?\s*$", b):
        return person("VALABILA_PANA_LA")
    if re.search(r"nascut[a]?(?:/a)?\s+la\s+data\s+de\s*$", b) or re.search(r"nascut[a]?(?:/a)?\s+la\s*$", b):
        return person("DATA_NASTERII")
    if re.search(r"\bin\s*$", b) and state.last_field == "DATA_NASTERII":
        return person("LOCUL_NASTERII")
    if re.search(r"domiciliat[a]?(?:/a)?\s+in\s*$", b) or re.search(r"domiciliul\s*(?:in)?\s*$", b):
        return person("ADRESA")
    if re.search(r"\bjud\.?\s*$", b) and state.last_scope == "person":
        return person("JUDET", "medium")
    if re.search(r"cetatenia\s*$", b) or re.search(r"cetatean\s*$", b):
        return person("CETATENIA", "medium")

    # — societate —
    # „Domeniul principal de activitate” = grupa CAEN (3 cifre, cu denumirea ei — derivată în aplicație din
    # codul principal), distinctă de „Activitatea principală” (clasa CAEN, 4 cifre — CAEN_1); unele acte cer
    # amândouă, una după alta.
    if re.search(r"domeniul principal de activitate este:?\s*$", b):
        return company("CAEN_DOMENIU")
    if re.search(r"activitatea principala este:?\s*$", b):
        return company("CAEN_1")
    if re.search(r"c\.?\s?u\.?\s?i\.?\s*:?$", b) or re.search(r"cod\s+fiscal\s*:?$", b):
        return company("SOCIETATE_CIF")
    if re.search(r"\bsediul(?:\s+social)?(?:\s+in)?\s*$", b) or re.search(r"\bsediu\s+social\s*(?:in)?\s*$", b):
        # dacă imediat după urmează un „jud. ……” separat, sediul nu mai include județul — altfel ar apărea de două ori
        if re.match(r"^[\s,]{0,6}jud\.?\b", a):
            return company("SOCIETATE_SEDIU_FARA_JUDET")
        return company("SOCIETATE_SEDIU")
    if re.search(r"\bjud\.?\s*$", b) and (state.last_field in ("SOCIETATE_SEDIU", "SOCIETATE_SEDIU_FARA_JUDET", "SOCIETATE_JUDET")
                                          or re.search(r"teritoriul\s+jud\.?\s*$", b)):
        return company("SOCIETATE_JUDET", "medium")
    if state.last_scope == "company" and state.last_field in ("SOCIETATE_SEDIU", "SOCIETATE_JUDET") and re.match(r"^,?\s*c\.?\s?u\.?\s?i", a):
        return company("SOCIETATE_NR_REG", "medium")
    if re.search(r"denumirea\s*$", b) or re.search(r"societatii\s*$", b) or re.match(r"^\s*(?:s\.?r\.?l\.?)?\s*(?:cu\s+sediul|,\s*cu\s+sediul)", a) \
            or re.search(r"(?:beneficiarii? reali? ai )?societatii$", bt):
        return company("SOCIETATE_DENUMIRE")

    # — capital și părți sociale —
    if re.match(r"^\s*lei\b", a) and re.search(r"(?:capital\w*\s+social.*valoare\s+de|valoare\s+de)\s*$", b) and not re.search(r"\baport", b):
        return company("CAPITAL_SOCIAL_TOTAL")
    if re.match(r"^\s*lei\b", a) and re.search(r"\baport\w*", b):
        return person("CAPITAL_SOCIAL")
    if re.match(r"^\s*part", a) and re.search(r"\bin\s*$", b):
        return company("PARTI_SOCIALE_TOTALE", "medium")
    if re.match(r"^\s*part", a) and re.search(r"revin\s*$", b):
        return person("PARTI_SOCIALE")

    # — persoane: nume —
    identity_follows = bool(re.match(r"^\s*,\s*(?:cetatean|cetatenia|domiciliat|nascut|identificat)", a))
    if identity_follows:                               # nume urmat de date de identitate = persoană definită aici
        return person("NUME_COMPLET", "high", new_person=True)
    if _MENTION_TRIGGER_RE.search(b):
        return person("NUME_COMPLET", "medium", mention=True)      # „asociatul ……”, „administratorul ……”: doar o mențiune
    if re.search(r"(?:\bsi\s*|,\s*)$", b) and state.last_field == "NUME_COMPLET":
        return person("NUME_COMPLET", "medium", mention=True)      # continuarea unei liste: „X, Y si Z”
    if re.match(r"^\s*_{5,}", after) and not before.strip():   # nume înaintea liniei de semnătură
        return person("NUME_COMPLET", "medium", mention=True)

    # — date —
    if re.search(r"\b(?:semnat\s+azi|incheiat\s+astazi|astazi|azi)\s*$", b) or re.search(r"\bnr\.?\s*\S*\s*din\s*$", b):
        return {"scope": "company", "field": "DATA_AZI", "confidence": "medium"}

    # — câmpuri manuale, dar cu etichetă curată (nu ghicită din ultimele cuvinte) —
    # „a câte …. pagini”: numărul de pagini al documentului nu se poate calcula automat (paginarea reală ține
    # de motorul de randare — Word, LibreOffice —, nu de conținutul .docx), deci rămâne completat manual.
    if re.search(r"\bcate\s*$", b) and re.match(r"^\s*\.?\s*pagini\b", a):
        return {"scope": "manual", "field": None, "confidence": "high", "label": "Număr de pagini"}
    # „redactat în …… exemplare” — numărul de exemplare originale, tipic scris literal (ex. „2 exemplare”) în
    # șablonul lawyer-ului, nu ca loc liber; dacă a fost înlocuit cu „……”, rămâne tot manual, dar etichetat clar.
    if re.match(r"^\s*exemplare\b", a):
        return {"scope": "manual", "field": None, "confidence": "high", "label": "Număr de exemplare"}
    return {"scope": "manual", "field": None, "confidence": "low"}


def _label_before(before: str) -> str:
    words = re.sub(r"[^\w\s]", " ", before).split()
    return " ".join(words[-4:]) if words else "valoare"


_MIN_DISPLAY_CONTEXT = 15   # caractere reale (fără spații) sub care se apelează la paragraful vecin, vezi mai jos
_NEARBY_PARAGRAPHS_LIMIT = 3   # câte paragrafe vecine NEVIDE se string, de fiecare parte, pentru contextul „larg”


def _nearby_paragraph_texts(all_texts: list[str], pi: int, direction: int, limit: int = _NEARBY_PARAGRAPHS_LIMIT) -> list[str]:
    """Textele (nevide) a până la `limit` paragrafe vecine lui `pi`, în direcția `direction` (+1 = după,
    -1 = înainte), SĂRIND paragrafele complet goale (spații vizuale între secțiuni, frecvente în actele
    reale) — cel mai apropiat paragraf nevid primul. Fără asta, un fallback care se uită doar la UN singur
    paragraf vecin rămâne fără niciun cuvânt de context dacă tocmai acela e gol (mai multe paragrafe goale
    la rând, ex. înainte de un titlu de secțiune — raportat de utilizator)."""
    out: list[str] = []
    j = pi + direction
    while 0 <= j < len(all_texts) and len(out) < limit:
        t = all_texts[j].strip()
        if t:
            out.append(t)
        j += direction
    return out


def _display_context(before: str, after: str, prev_pars: list[str], next_pars: list[str]) -> tuple[str, str, str, str]:
    """(before, after, before_larg, after_larg) — primele două, de AFIȘAT implicit în interfață: dacă locul
    liber e (aproape) singur în paragraful lui („……” pe propriul rând, sub un titlu, ex. denumirea firmei),
    textul din ACEST paragraf nu ajunge ca reper vizual pentru om („... ... ...”, fără niciun cuvânt real) —
    se completează cu cel mai apropiat paragraf vecin NEVID (poate sări peste unul sau mai multe goale — vezi
    _nearby_paragraph_texts), separat printr-un „ ¶ ” (ruptura de paragraf). Ultimele două — context mult mai
    larg (până la _NEARBY_PARAGRAPHS_LIMIT paragrafe de fiecare parte), arătat DOAR la cerere (buton „arată
    mai mult”, vezi frontend), pentru orice loc liber, nu doar cele cu context sărac — cerere utilizator.
    Niciuna din ele nu intră la POTRIVIREA regulilor (asta rămâne pe fereastra strict din acest paragraf —
    vezi analyze())."""
    db = before
    if len(before.strip()) < _MIN_DISPLAY_CONTEXT and prev_pars:
        db = prev_pars[0][-90:].strip() + " ¶ " + before
    da = after
    if len(after.strip()) < _MIN_DISPLAY_CONTEXT and next_pars:
        da = after + " ¶ " + next_pars[0][:90].strip()
    wide_before = (" ¶ ".join(reversed(prev_pars)) + " ¶ " + before) if prev_pars else before
    wide_after = (after + " ¶ " + " ¶ ".join(next_pars)) if next_pars else after
    return db, da, wide_before[-500:], wide_after[:500]


def _finalize_blank(out: list[dict], s: dict, pi: int, start: int, end: int, before: str, after: str,
                    manual_label: str | None = None, prev_pars: list[str] | None = None,
                    next_pars: list[str] | None = None) -> None:
    """Completează poziția + eticheta + eticheta finală (tag) unui loc liber găsit și îl adaugă la `out` —
    pasul comun celor două ramuri din analyze() (cu și fără indiciu în paranteză, care doar pregătesc `s` diferit
    până aici). `manual_label`, dat doar de ramura cu indiciu, e eticheta explicită scrisă de autor
    („(Nume Asociat)” → „Nume Asociat”); fără el, un câmp manual ia eticheta deja propusă de suggest() sau,
    în lipsă, ultimele cuvinte dinaintea locului liber (vezi _label_before). `prev_pars`/`next_pars` (paragrafele
    vecine) intră DOAR în ce se afișează (vezi _display_context), niciodată în potrivirea regulilor de mai sus."""
    db, da, wb, wa = _display_context(before, after, prev_pars or [], next_pars or [])
    s.update({"id": len(out), "paragraph": pi, "start": start, "end": end,
              "before": db, "after": da, "before_wide": wb, "after_wide": wa})
    if s["scope"] == "manual":
        s["label"] = manual_label if manual_label is not None else (s.get("label") or _label_before(before))
        s["tag"] = manual_tag(s["label"])
    elif s["scope"] == "company":
        s["label"] = COMPANY_FIELDS[s["field"]]
        s["tag"] = "{{" + s["field"] + "}}"
    else:
        s["label"] = f'{PERSON_FIELDS[s["field"]]} ({"administrator" if s["role"] == "ADMINISTRATOR" else "asociat"} {s["person"]})'
        s["tag"] = person_tag(s["role"], s["person"], s["field"])
    out.append(s)


_LIST_FIELD = {"ASOCIAT": "ASOCIATI_LISTA", "ADMINISTRATOR": "ADMINISTRATORI_LISTA"}
_CONNECTOR_RE = re.compile(r"^[\s,]*(?:si|și)?[\s,]*$", re.IGNORECASE)


# ── Obiectul de activitate (CAEN): uneori rămâne needitat la pregătirea documentului — nu are „……”, ci
# codurile și denumirile activităților firmei-exemplu din care a pornit șablonul („Domeniul principal de
# activitate este: 953 Repararea ...”, „Activitatea principală este: 9531 Repararea ...”, urmate de o listă de
# activități secundare, câte una pe rând, fără nicio frază introductivă proprie: „4672 Comerț cu ridicata
# ...”). Fără loc liber de găsit, secțiunea asta se recunoaște din formulare + cod CAEN (4 cifre), nu din BLANK_RE.
_CAEN_DIACRITICS = str.maketrans("ăâîșşțţĂÂÎȘŞȚŢ", "aaissttAAISSTT")


def _ascii(s: str) -> str:
    """Variantă fără diacritice, cu aceeași lungime ca originalul — spre deosebire de normalizarea Unicode NFD
    din variants._norm (schimbă lungimea), pozițiile găsite prin regex pe ea rămân valide pe textul original."""
    return s.translate(_CAEN_DIACRITICS).lower()


def _par_text(par) -> str:
    return "".join(r.text for r in par.runs)


_CAEN_DOMENIU_RE = re.compile(r"domeniul principal de activitate este:?\s*(\d{3,4})\s+(.+?)\s*(?:;|-{2,}|$)")
_CAEN_PRINCIPAL_RE = re.compile(r"activitatea principala este:?\s*(\d{4})\s+(.+?)\s*(?:;|-{2,}|$)")
_CAEN_LIST_TRIGGER_RE = re.compile(r"urmatoarele activitat")
_CAEN_LINE_RE = re.compile(r"^\s*(\d{4})\s+(.+?)\s*-{0,}\s*$")


def _find_caen_singles(text: str) -> list[dict]:
    """„Domeniul principal de activitate este: 953 ...” / „Activitatea principală este: 9531 ...” — codul și
    denumirea (nu fraza dinainte) devin locul liber: domeniul (grupa CAEN, 3 cifre) → {{CAEN_DOMENIU}},
    activitatea principală (clasa CAEN, 4 cifre) → {{CAEN_1}} — aplicația derivă automat grupa din clasă
    (vezi frontend/src/data/caenGrupe.ts), deci cele două rămân consistente fără completare separată."""
    norm = _ascii(text)
    out = []
    for rx, field in ((_CAEN_DOMENIU_RE, "CAEN_DOMENIU"), (_CAEN_PRINCIPAL_RE, "CAEN_1")):
        for m in rx.finditer(norm):
            out.append({"scope": "company", "field": field, "confidence": "high", "start": m.start(1), "end": m.end(2)})
    return out


def _find_caen_secondary_cluster(paragraphs: list, pi: int) -> list[int]:
    """Dacă paragraful `pi` e declanșatorul listei de activități secundare („...următoarele activități:”),
    întoarce indicii paragrafelor consecutive de după el care sunt doar „<cod CAEN> <denumire>”, fără nimic
    altceva (o listă de-a lungul mai multor paragrafe, nu în aceeași frază) — sau listă goală, altfel."""
    if not _CAEN_LIST_TRIGGER_RE.search(_ascii(_par_text(paragraphs[pi]))):
        return []
    lines, j = [], pi + 1
    while j < len(paragraphs) and _CAEN_LINE_RE.match(_par_text(paragraphs[j])):
        lines.append(j)
        j += 1
    return lines


# ── Indiciu explicit din paranteze: „…(CAEN PRINCIPAL)……” — utilizatorul scrie chiar el, lipit de locul
# liber, ce reprezintă acesta, când nu are încredere că formularea din jur e destul de clară pentru
# recunoașterea automată. Are prioritate maximă, înaintea oricărei ghiciri din context (vezi suggest()).
_HINT_RE = re.compile(r"\s*\(([^()]{1,80})\)")
_HINT_ALIASES = {
    "caen principal": ("company", "CAEN_1"),
    "caen secundar": ("company", "CAEN"),
    "caen secundara": ("company", "CAEN"),
    "caen secundare": ("company", "CAEN"),
    "caen secundari": ("company", "CAEN"),
    "caen domeniu": ("company", "CAEN_DOMENIU"),
    "domeniu caen": ("company", "CAEN_DOMENIU"),
    "domeniul caen": ("company", "CAEN_DOMENIU"),
    "domeniu principal": ("company", "CAEN_DOMENIU"),
    "domeniul principal": ("company", "CAEN_DOMENIU"),
}


def _resolve_hint(hint: str) -> tuple[str, str] | None:
    """(scope, field) pentru un indiciu din paranteze, sau None dacă nu se recunoaște — caz în care indiciul
    devine el însuși eticheta unui câmp manual (vezi analyze()), nu se ghicește nimic. Acceptă fie un alias
    cunoscut (tabelul de mai sus), fie chiar numele câmpului intern, scris direct (ex. „(SOCIETATE_SEDIU)”)."""
    norm = re.sub(r"[^a-z0-9]+", " ", _ascii(hint)).strip()
    if not norm:
        return None
    if norm in _HINT_ALIASES:
        return _HINT_ALIASES[norm]
    key = norm.upper().replace(" ", "_")
    if key in COMPANY_FIELDS:
        return ("company", key)
    if key in PERSON_FIELDS:
        return ("person", key)
    return None


def analyze(docx_bytes: bytes) -> list[dict]:
    """Lista locurilor libere, în ordinea documentului, cu context și propunere."""
    doc = Document(io.BytesIO(docx_bytes))
    pars = _paragraphs(doc)
    all_par_texts = [_par_text(p) for p in pars]   # o singură trecere — reutilizat mai jos pentru context (vezi _nearby_paragraph_texts)
    out: list[dict] = []
    sig_counts: dict[str, int] = {}    # poziția (ASOCIAT_N/ADMINISTRATOR_N) următoarei linii de semnătură, per rol
    active_sig_role: str | None = None  # rolul secțiunii de semnătură „curente” — vezi _signature_role_for_heading
    for pi, par in enumerate(pars):
        text = all_par_texts[pi]
        # Doar pentru afișare (vezi _display_context) — un loc liber (aproape) singur în paragraful lui
        # (ex. denumirea firmei, sub titlu) altfel n-ar avea niciun cuvânt real de context în interfață. Poate
        # sări peste mai multe paragrafe goale consecutive (vezi _nearby_paragraph_texts), nu doar unul.
        prev_pars = _nearby_paragraph_texts(all_par_texts, pi, -1)
        next_pars = _nearby_paragraph_texts(all_par_texts, pi, 1)

        # Titlu de secțiune de semnătură („SEMNĂTURA ASOCIAT/ASOCIAȚI”, „SEMNATURILE,”) sau doar eticheta unui
        # rol, singură pe rând („Administrator”) — activează/schimbă rolul curent pentru liniile de semnătură
        # care urmează. Un paragraf cu conținut obișnuit, substanțial, înseamnă că am ieșit din zona de
        # semnătură (paragrafele scurte/goale dintre titlu și linia efectivă de semnat nu-l dezactivează).
        heading_role = _signature_role_for_heading(_ascii(text).strip(" \t.,:;-"))
        if heading_role is not None:
            active_sig_role = heading_role
        elif len(text.split()) > 8:
            active_sig_role = None

        # obiectul de activitate (CAEN), needitat de la firma-exemplu — nu are „……”, se recunoaște din formulare,
        # nu din BLANK_RE (vezi _find_caen_singles/_find_caen_secondary_cluster mai sus)
        for cs in _find_caen_singles(text):
            cs.update({"id": len(out), "paragraph": pi, "before": text[max(0, cs["start"] - 70):cs["start"]],
                       "after": text[cs["end"]:cs["end"] + 40], "label": COMPANY_FIELDS[cs["field"]], "tag": "{{" + cs["field"] + "}}"})
            out.append(cs)
        for lj in _find_caen_secondary_cluster(pars, pi):
            line_text = _par_text(pars[lj])
            lm = _CAEN_LINE_RE.match(line_text)
            out.append({
                "id": len(out), "paragraph": lj, "scope": "company", "field": "CAEN", "role": "CAEN", "confidence": "high",
                "start": lm.start(1), "end": lm.end(2), "before": "", "after": line_text[lm.end(2):lm.end(2) + 40],
                "label": COMPANY_FIELDS["CAEN"], "tag": "{{CAEN}}",
            })

        # liniile lungi de „_____” sunt spații de semnătură, nu locuri de completat
        matches = [m for m in BLANK_RE.finditer(text) if not (set(m.group()) == {"_"} and len(m.group()) >= 10)]
        if not matches:
            continue

        # Calitatea scrisă DUPĂ identitatea completă a persoanei („……, domiciliat…, CNP…, în calitate de
        # COMODANT”) — actele reale o scriu des la finalul clauzei, nu înainte de nume (unde se uită
        # _role_for). Pre-scanăm locurile care par să înceapă o clauză de persoană — nume urmat de identitate
        # sau precedat de un cuvânt de calitate/adresare, aceleași condiții ca în suggest(), dar fără să
        # atingă `state` (care ține numerotarea reală, mai jos) — ca să delimităm fiecare clauză (de la
        # începutul ei până la următoarea/finalul paragrafului) și să căutăm „în calitate de X” în ea.
        clause_starts: list[int] = []
        for cm in matches:
            cb = _norm(text[max(0, cm.start() - 110):cm.start()])[-90:]
            ca = _norm(text[cm.end():cm.end() + 50])[:50]
            # „identificat/domiciliat/…” DUPĂ locul liber înseamnă „e un nume” doar dacă ÎNAINTE de el nu e
            # deja text — altfel orice câmp urmat mai departe de „identificată cu CI seria…” (ex. județul,
            # în „jud. ……, identificată cu CI…”) ar fi luat greșit drept începutul unei clauze noi.
            starts_by_identity = not cb.strip(" ,") and bool(
                re.match(r"^\s*,\s*(?:cetatean|cetatenia|domiciliat|nascut|identificat)", ca))
            if starts_by_identity or _MENTION_TRIGGER_RE.search(cb):
                clause_starts.append(cm.start())
        clause_role: dict[int, str] = {}
        for ci, cstart in enumerate(clause_starts):
            cend = clause_starts[ci + 1] if ci + 1 < len(clause_starts) else len(text)
            explicit = _explicit_role(text[cstart:cend])
            if explicit:
                clause_role[cstart] = explicit

        def _role_override(pos: int) -> str | None:
            """Rolul explicit al clauzei căreia îi aparține poziția `pos` (vezi mai sus), sau None dacă acea
            clauză n-are nicio calitate declarată explicit."""
            current = max((cs for cs in clause_starts if cs <= pos), default=None)
            return clause_role.get(current) if current is not None else None

        state = _State()
        list_entry: dict | None = None     # entry-ul „ASOCIATI_LISTA”/„ADMINISTRATORI_LISTA” aflat în curs de extindere
        list_role: str | None = None       # rolul lui list_entry — ținut separat, fiindcă odată convertit entry-ul nu mai are „role”
        consumed_hint_blanks: set[int] = set()   # indicele (în `matches`) al blank-ului al 2-lea dintr-un „…(indiciu)……” deja înghițit
        for k, m in enumerate(matches):
            if k in consumed_hint_blanks:
                continue
            blank_end = _extend_past_artifact_dots(text, m.end())    # vezi _extend_past_artifact_dots
            # Fereastră largă (mult peste ce citește suggest() pentru potrivire, care își taie singur cât are
            # nevoie mai jos) — asta e și ce vede utilizatorul ca „before”/„after” la un câmp needeslușit; prea
            # îngustă, arăta adesea doar alte locuri libere vecine („…… ……”), fără niciun cuvânt real de context.
            before = text[max(0, m.start() - 200):m.start()]
            after = text[blank_end:blank_end + 110]

            # indiciu explicit lipit de locul liber, ex. „…(CAEN PRINCIPAL)……” — prioritate maximă, înaintea
            # oricărei ghiciri din context (vezi _resolve_hint); dacă mai urmează imediat un al doilea loc
            # liber (doar spații între ele), cele două + indiciul devin O SINGURĂ etichetă, nu două separate.
            hint_m = _HINT_RE.match(text, blank_end)
            if hint_m:
                end = hint_m.end()
                nxt = k + 1
                if nxt < len(matches) and not text[end:matches[nxt].start()].strip():
                    end = _extend_past_artifact_dots(text, matches[nxt].end())
                    consumed_hint_blanks.add(nxt)
                resolved = _resolve_hint(hint_m.group(1))
                if resolved:
                    scope, field = resolved
                    if scope == "company":
                        s = {"scope": "company", "field": field, "confidence": "high"}
                        if field == "CAEN":
                            s["role"] = "CAEN"
                    else:
                        # Ca la câmpurile atașate din person() de mai sus — persoana „curentă” din paragraf,
                        # nu neapărat cea „definită complet” (vezi _State.last_person_n). Un semnal real
                        # ÎNAINTE de locul liber are prioritate pe calitatea găsită DUPĂ (vezi _role_for_hint).
                        role = _role_for_hint(text[:m.start()]) or _role_override(m.start()) or "ASOCIAT"
                        s = {"scope": "person", "field": field, "role": role,
                             "person": state.last_person_n or 1, "confidence": "high"}
                    state.last_field, state.last_scope = field, scope
                elif active_sig_role is not None and _looks_like_signature_remainder(before + text[end:]):
                    # indiciul nu s-a recunoscut ca un câmp anume (ex. „(Nume Asociat)”), dar locul liber
                    # (+ indiciu) e practic tot ce are paragraful, lângă titlul unei secțiuni de semnătură —
                    # tratat ca linie de semnătură, nu ca un câmp manual generic (vezi _signature_role_for_heading).
                    sig_counts[active_sig_role] = sig_counts.get(active_sig_role, 0) + 1
                    s = {"scope": "person", "field": "NUME_COMPLET", "role": active_sig_role,
                         "person": sig_counts[active_sig_role], "confidence": "high", "is_signature_line": True}
                else:
                    s = {"scope": "manual", "field": None, "confidence": "high", "_hint_label": hint_m.group(1).strip()}
                _finalize_blank(out, s, pi, m.start(), end, before, text[end:end + 110],
                                manual_label=(s.pop("_hint_label") if s["scope"] == "manual" else None),
                                prev_pars=prev_pars, next_pars=next_pars)
                list_entry, list_role = None, None
                continue

            alone_here = text.strip(" -\t") == text[m.start():blank_end]
            s = suggest(before, after, text[:m.start()], state, pi, alone=alone_here)
            if s.get("scope") == "person" and _role_for_hint(text[:m.start()]) is None:
                # Calitatea scrisă oriunde în clauza persoanei (vezi clause_role mai sus) — suggest()/_role_for
                # văd doar ce e ÎNAINTE de locul liber, dar actele reale scriu des calitatea la finalul
                # clauzei; suprascriem DOAR când textul dinainte n-avea deja un semnal real (vezi _role_for_hint
                # — un semnal găsit înainte, ex. „administrator”, nu trebuie suprascris de o calitate care,
                # mai departe, descrie de fapt PARTEA/firma, nu persoana).
                override = _role_override(m.start())
                if override:
                    s["role"] = override
            is_mention = s["scope"] == "person" and s["field"] == "NUME_COMPLET" and s["confidence"] == "medium"
            # Linie de semnătură — fie tiparul vechi (nume urmat de „_____” pe același rând), fie locul liber
            # singur pe rând, lângă titlul unei secțiuni de semnătură (vezi _signature_role_for_heading).
            # Rolul din titlu (`active_sig_role`), când există, are prioritate pe rolul ghicit de suggest()
            # (care, pentru un rând gol fără niciun cuvânt înainte, oricum nu are de unde să-l deducă).
            sig_role = None
            if is_mention and re.match(r"^\s*_{5,}", after) and not before.strip():
                sig_role = active_sig_role if active_sig_role is not None else s["role"]
            elif active_sig_role is not None and alone_here:
                sig_role = active_sig_role
            if sig_role is not None:
                sig_counts[sig_role] = sig_counts.get(sig_role, 0) + 1   # câte o poziție per rol, nu una globală
                s["scope"], s["field"], s["role"], s["person"] = "person", "NUME_COMPLET", sig_role, sig_counts[sig_role]
                s["confidence"] = "high" if active_sig_role is not None else "medium"
                s["is_signature_line"] = True
                is_mention = False                         # rămân individuale (câte o linie de semnat per persoană)
            # A doua (sau a N-a) mențiune consecutivă a aceluiași rol, „…… și ……”: în loc de poziții fixe
            # (ASOCIAT_1, ASOCIAT_2 — nu au loc pentru o a treia persoană dacă apare), entry-ul anterior devine un
            # singur câmp scalabil (ASOCIATI_LISTA), care merge la orice număr — vezi lib/placeholders.ts (joinNames).
            # O mențiune SINGURĂ (fără alta lângă ea) rămâne individuală, ca înainte — ar putea desemna o persoană
            # anume (ex. „administratorul ……” dintr-o singură mențiune), nu neapărat lista completă.
            # `list_role in _LIST_FIELD` — doar Asociat/Administrator au azi un câmp „…LISTA” dedicat, cu
            # sprijin în frontend (joinNames); un rol nou (comodant, reprezentant legal…) rămâne cu mențiuni
            # individuale, numerotate, mai degrabă decât să inventăm o etichetă „…LISTA” pe care nimic n-o completează.
            if is_mention and list_entry is not None and list_role == s["role"] and list_role in _LIST_FIELD \
                    and _CONNECTOR_RE.match(text[list_entry["end"]:m.start()]):
                list_entry["end"] = blank_end
                list_entry["after"] = after
                if list_entry["scope"] != "company":
                    field = _LIST_FIELD[list_role]
                    list_entry.update(scope="company", field=field, role=None, person=None,
                                      label=COMPANY_FIELDS[field], tag="{{" + field + "}}")
                continue
            _finalize_blank(out, s, pi, m.start(), blank_end, before, after, prev_pars=prev_pars, next_pars=next_pars)
            list_entry, list_role = (s, s["role"]) if is_mention else (None, None)
    return out


# ── Blocuri repetitive: câte o persoană cu aceeași structură, repetată de câte ori e nevoie ──────────────────

_ROLE_PLURAL = {"ASOCIAT": "ASOCIATI", "ADMINISTRATOR": "ADMINISTRATORI", "CAEN": "CAEN_SECUNDARE"}
_ORDINAL_RE = re.compile(r"^(\s*)(\d+)(\.\s*)")

# „asociatul/asociații”, „administratorul/administratorii”, dar la fel de bine „comodantul/comodatarul”,
# „locatorul/locatarul” — ORICE cuvânt din _ROLE_KEYWORDS, nu doar cele două roluri din actul constitutiv (cu
# variații: verb schimbat — „este asociatul/sunt asociații” —, sau typo real „adminstratorul”). Construit din
# _ROLE_KEYWORDS (ca _ROLE_LABEL_RE mai sus), cu o singură excepție: ADMINISTRATOR păstrează fragmentul mai
# tolerant `adm\w*strat\w*` (prinde și typo-ul „adminstrator”, căruia îi lipsește un „i” — găsit într-un
# document real; `administr\w*` din _ROLE_KEYWORDS nu-l prinde, fiindcă typo-ul rupe exact acel prefix).
_PLURAL_MARKER_ALT = "|".join(
    (r"adm\w*strat\w*" if role == "ADMINISTRATOR" else kw.replace(" ", r"\s+") + r"\w*")
    for kw, role in _ROLE_KEYWORDS
)
_PLURAL_MARKER_RE = re.compile(rf"({_PLURAL_MARKER_ALT})\s*/\s*(?:\w+\s+)?({_PLURAL_MARKER_ALT})")


def _has_plural_marker(text: str) -> bool:
    return bool(_PLURAL_MARKER_RE.search(_ascii(text)))


def _paragraph_has_plural_marker(paragraphs: list, pi: int, marker_start: int) -> bool:
    """Marcajul poate fi în ACELAȘI paragraf, înaintea persoanei (fraza obișnuită), sau — cum apare des în
    documentul real — într-un paragraf Word separat, chiar înaintea celui cu datele persoanei (o ruptură de
    paragraf pe care autorul n-a intenționat-o, dar care există în fișier)."""
    if _has_plural_marker(_par_text(paragraphs[pi])[:marker_start]):
        return True
    return pi > 0 and _has_plural_marker(_par_text(paragraphs[pi - 1])[-160:])


def detect_groups(blanks: list[dict], docx_bytes: bytes | None = None) -> list[dict]:
    """Grupuri candidate de bloc repetitiv: mai multe persoane cu exact aceeași structură de câmpuri, fie în
    aceeași frază ([]„X, cetățenia…, CNP … si Y, cetățenia…, CNP …”] — kind='inline'), fie în paragrafe separate
    consecutive (kind='paragraph'), fie o SINGURĂ persoană scrisă ca exemplu, dar cu marcaj explicit de plural
    alături („asociatul/asociații ……, cetățean…” — vezi _paragraph_has_plural_marker; are nevoie de `docx_bytes`,
    ca să poată citi și paragraful dinainte). Nu se schimbă nimic aici — doar se propun; utilizatorul alege în
    interfață dacă se repetă automat (scalează la orice număr) sau rămân poziții fixe (vezi apply())."""
    groups: list[dict] = []
    by_par: dict[int, list[dict]] = {}
    for b in blanks:
        by_par.setdefault(b["paragraph"], []).append(b)
    for items in by_par.values():
        by_par[items[0]["paragraph"]] = sorted(items, key=lambda x: x["start"])

    def _is_name(b: dict) -> bool:
        return b["scope"] == "person" and b["field"] == "NUME_COMPLET"

    def _clause_marker_indices(person_items: list[dict]) -> list[int]:
        """Indicii din `person_items` care ANCOREAZĂ o clauză (o persoană nouă): un NUME_COMPLET urmat de cel
        puțin un ALT câmp al aceleiași persoane, înainte de următorul NUME_COMPLET — indiferent dacă numele a
        fost introdus complet („……, cetățean…, CNP…”, confidence „high”) sau doar amintit („Asociatul ……
        contribuie cu un aport de …… lei”, confidence „medium”; vezi suggest()). Generic — orice câmp poate fi
        cel „atașat” (aport, cotă, procent…), nu doar identitatea completă — ca detectarea să prindă tipare
        similare din alte șabloane, nu doar cazul CNP/domiciliat. O mențiune GOALĂ (numele, fără nimic
        atașat) nu ancorează nimic — rămâne o simplă mențiune, scalabilă separat ca listă de nume (vezi
        ASOCIATI_LISTA în analyze()), nu ca bloc repetitiv cu clauze."""
        name_idx = [i for i, b in enumerate(person_items) if _is_name(b)]
        return [i for k, i in enumerate(name_idx)
                if (name_idx[k + 1] if k + 1 < len(name_idx) else len(person_items)) > i + 1]

    # kind='inline' — mai multe persoane cu clauze în ACELAȘI paragraf (unite prin „si”/„iar”), fie definite
    # complet, fie doar amintite dar cu un câmp atașat (vezi _clause_marker_indices). ADRESA e tratată ca
    # opțională la potrivirea formei: în documente reale, o persoană e uneori scrisă „domiciliat/ă în jud.……”
    # (doar județ, fără localitate separată) și alta „domiciliat/ă în ……, jud. ……” (localitate + județ) —
    # aceeași clauză, doar o mențiune mai scurtă a domiciliului. Fără relaxarea asta, o asemenea diferență
    # reală blochează gruparea și persoanele suplimentare dispar din acel paragraf.
    for pi, items in sorted(by_par.items()):
        person_items = [b for b in items if b["scope"] == "person"]
        markers = _clause_marker_indices(person_items)
        if len(markers) < 2:
            continue
        bounds = markers + [len(person_items)]
        clauses_raw = [person_items[bounds[k]:bounds[k + 1]] for k in range(len(markers))]
        role = clauses_raw[0][0]["role"]
        if any(not clause or any(c["role"] != role for c in clause) for clause in clauses_raw):
            continue
        shapes = [tuple(c["field"] for c in clause) for clause in clauses_raw]
        norm_shapes = [tuple(f for f in s if f != "ADRESA") for s in shapes]
        if len(set(norm_shapes)) != 1:
            continue
        clauses = clauses_raw
        template_idx = max(range(len(shapes)), key=lambda i: len(shapes[i]))    # clauza cea mai completă, ca șablon
        blank_ids = [b["id"] for c in clauses for b in c]
        groups.append({
            "id": len(groups), "kind": "inline", "paragraph": pi, "role": role, "count": len(clauses),
            "blank_ids": blank_ids, "template_blank_ids": [b["id"] for b in clauses[template_idx]],
            "label": f'{"administratori" if role == "ADMINISTRATOR" else "asociați"} — {len(clauses)} persoane găsite în același paragraf',
        })

    # kind='paragraph' — paragrafe separate consecutive, fiecare cu exact o persoană (numele — definit complet
    # sau doar amintit —, urmat de restul câmpurilor ei), aceeași structură
    solo: list[tuple[int, str, tuple, list[dict]]] = []
    for pi, items in sorted(by_par.items()):
        person_items = [b for b in items if b["scope"] == "person"]
        if len(person_items) >= 2 and _is_name(person_items[0]) and sum(_is_name(b) for b in person_items) == 1:
            solo.append((pi, person_items[0]["role"], tuple(b["field"] for b in person_items), person_items))
    i = 0
    while i < len(solo):
        j = i
        pi0, role0, shape0, _ = solo[i]
        cluster = [solo[i]]
        while j + 1 < len(solo) and solo[j + 1][0] == cluster[-1][0] + 1 and solo[j + 1][1] == role0 and solo[j + 1][2] == shape0:
            cluster.append(solo[j + 1])
            j += 1
        if len(cluster) >= 2:
            groups.append({
                "id": len(groups), "kind": "paragraph", "paragraphs": [c[0] for c in cluster], "role": role0,
                "count": len(cluster), "blank_ids": [b["id"] for c in cluster for b in c[3]],
                "template_blank_ids": [b["id"] for b in cluster[0][3]],
                "label": f'{"administratori" if role0 == "ADMINISTRATOR" else "asociați"} — {len(cluster)} paragrafe consecutive, aceeași structură',
            })
        i = j + 1

    # kind='paragraph', role='CAEN' — activități CAEN secundare needitate de la firma-exemplu, câte un cod pe
    # rând (fără „……”, vezi _find_caen_secondary_cluster); spre deosebire de persoane, un cod pe rând e deja
    # o „clauză” completă, deci nu mai e nevoie de un prag minim de 2 — chiar și un singur cod găsit merită
    # convertit în bloc repetitiv, ca lista reală a clientului (oricâte activități are) să încapă.
    caen_lines = sorted((b for items in by_par.values() for b in items if b.get("role") == "CAEN"), key=lambda b: b["paragraph"])
    i = 0
    while i < len(caen_lines):
        j = i
        cluster = [caen_lines[i]]
        while j + 1 < len(caen_lines) and caen_lines[j + 1]["paragraph"] == cluster[-1]["paragraph"] + 1:
            cluster.append(caen_lines[j + 1])
            j += 1
        cod_word = "cod găsit" if len(cluster) == 1 else "coduri găsite"
        groups.append({
            "id": len(groups), "kind": "paragraph", "paragraphs": [b["paragraph"] for b in cluster], "role": "CAEN",
            "count": len(cluster), "blank_ids": [b["id"] for b in cluster], "template_blank_ids": [cluster[0]["id"]],
            "label": f'activități CAEN secundare — {len(cluster)} {cod_word}, câte unul pe rând',
        })
        i = j + 1

    # kind='paragraph' — liniile de semnătură (vezi analyze(): sig_counts/_signature_role_for_heading),
    # grupate pe rol — NU pe adiacența paragrafelor ca la „solo” mai sus: între liniile de semnat există des
    # paragrafe goale (spațiu vizual pentru semnătura olografă), care ar rupe o cerință de „paragraf imediat
    # următor”. Fără gruparea asta, șablonul are loc doar pentru câte linii de semnătură a scris autorul (ex.
    # 2), indiferent câți asociați/administratori are efectiv clientul (ex. 3) — al treilea n-ar avea unde
    # semna. Chiar și o singură linie găsită merită bloc repetitiv (ca la CAEN mai sus), pentru același motiv.
    sig_lines = sorted((b for items in by_par.values() for b in items if b.get("is_signature_line")),
                       key=lambda b: (b["role"], b["paragraph"]))
    for role, role_lines in itertools.groupby(sig_lines, key=lambda b: b["role"]):
        cluster = list(role_lines)
        if role == "ADMINISTRATOR":
            role_word = "administratori"
        elif role == "ASOCIAT":
            role_word = "asociați"
        else:
            role_word = role.replace("_", " ").lower()
        groups.append({
            "id": len(groups), "kind": "paragraph", "paragraphs": [b["paragraph"] for b in cluster], "role": role,
            "count": len(cluster), "blank_ids": [b["id"] for b in cluster], "template_blank_ids": [cluster[0]["id"]],
            "label": f'{role_word} — {len(cluster)} {"linie" if len(cluster) == 1 else "linii"} de semnătură',
        })

    # kind='inline', o SINGURĂ clauză — „asociatul/asociații ……, cetățean…” sau „asociatul/asociații ……
    # contribuie cu un aport de …… lei” — marcaj explicit de plural, chiar dacă e scrisă o singură persoană
    # ca exemplu (vezi _paragraph_has_plural_marker mai sus). Se oferă exact ca la 2+ clauze: „repeat”
    # scalează la orice număr, „fixed” rămâne o singură poziție.
    if docx_bytes is not None:
        pars = _paragraphs(Document(io.BytesIO(docx_bytes)))
        already_grouped = {g["paragraph"] for g in groups if g["kind"] == "inline"} | \
            {pi for g in groups if g["kind"] == "paragraph" for pi in g["paragraphs"]}
        for pi, items in sorted(by_par.items()):
            if pi in already_grouped or pi >= len(pars):
                continue
            person_items = [b for b in items if b["scope"] == "person"]
            markers = _clause_marker_indices(person_items)
            if len(markers) != 1:
                continue
            clause = person_items[markers[0]:]
            if not _paragraph_has_plural_marker(pars, pi, clause[0]["start"]):
                continue
            role = clause[0]["role"]
            groups.append({
                "id": len(groups), "kind": "inline", "paragraph": pi, "role": role, "count": 1,
                "blank_ids": [b["id"] for b in clause], "template_blank_ids": [b["id"] for b in clause],
                "label": f'{"administrator" if role == "ADMINISTRATOR" else "asociat"} — marcat cu '
                         f'„{"administratorul/administratorii" if role == "ADMINISTRATOR" else "asociatul/asociații"}”, '
                         f'o singură persoană scrisă ca exemplu',
            })
    return groups


def generic_person_tag(field: str) -> str:
    """Eticheta nenumerotată pentru un câmp de persoană dintr-un bloc {{#ASOCIATI}}/{{#ADMINISTRATORI}} — aceleași
    nume ca în lib/placeholders.ts (persoanaToSingularMap)."""
    if field == "NUME_COMPLET":
        return "{{NUME}} {{PRENUME}}"
    if field == "CAPITAL_SOCIAL":
        return "{{CAPITAL_SOCIAL}}"
    if field == "PARTI_SOCIALE":
        return "{{PARTI_SOCIALE}}"
    return "{{" + field + "}}"


def _dominant_rpr(paragraph: Paragraph):
    """Formatarea (rPr — font, mărime, bold…) a fugii (run) cu cel mai mult text din paragraf: reprezintă fontul
    de bază al paragrafului (Arial, Times New Roman, ce-o fi în șablon), nu neapărat al locului liber „……” în
    sine, care câteodată are formatare proprie (ex. subliniat)."""
    runs = [r for r in paragraph.runs if r.text]
    if not runs:
        return None
    dominant = max(runs, key=lambda r: len(r.text))
    rpr = dominant._r.find(qn("w:rPr"))
    return copy.deepcopy(rpr) if rpr is not None else None


def _rpr_at(paragraph: Paragraph, pos: int):
    """rPr al fugii care acoperă poziția `pos` din textul concatenat al paragrafului — formatarea PROPRIE a
    unui loc anume (ex. „……” scris cu sublinire, diferit de restul frazei), nu formatarea dominantă a
    întregului paragraf (vezi _dominant_rpr). None dacă poziția cade în afara textului sau fuga n-are rPr."""
    cursor = 0
    for r in paragraph.runs:
        end = cursor + len(r.text)
        if cursor <= pos < end:
            rpr = r._r.find(qn("w:rPr"))
            return copy.deepcopy(rpr) if rpr is not None else None
        cursor = end
    return None


def _insert_paragraph_before(anchor: Paragraph, content, rpr=None) -> Paragraph:
    """Paragraf nou, gol, inserat imediat înaintea lui `anchor`, cu aceeași aliniere/indentare (pPr). `content`
    e fie un șir simplu (o singură fugă, cu formatarea `rpr` — vezi _dominant_rpr), fie o listă de segmente
    (text, rPr_propriu_sau_None) — vezi _local_replacements: fiecare segment devine propria lui fugă, cu
    formatarea EI dacă are una (ex. locul liber era subliniat), altfel cea dominantă (`rpr`) — altfel Word
    afișează fugile noi cu fontul implicit (adesea Times New Roman), diferit de restul șablonului (ex. Arial)."""
    new_p = OxmlElement("w:p")
    if anchor._p.pPr is not None:
        new_p.append(copy.deepcopy(anchor._p.pPr))
    anchor._p.addprevious(new_p)
    new_par = Paragraph(new_p, anchor._parent)
    segments = content if isinstance(content, list) else ([(content, None)] if content else [])
    for text, seg_rpr in segments:
        if not text:
            continue
        run = new_par.add_run(text)
        chosen_rpr = seg_rpr if seg_rpr is not None else rpr
        if chosen_rpr is not None:
            run._r.insert(0, copy.deepcopy(chosen_rpr))
    return new_par


def _local_replacements(text: str, blanks: list[dict], offset: int, tag_for, rpr_for=None) -> list[tuple[str, object]]:
    """Aplică `tag_for(blank)` (poate întoarce None = neschimbat) pe `text`, pentru blank-urile din `blanks`
    (coordonate absolute, decalate cu -offset). Întoarce segmente (text, rPr_sau_None) pentru
    _insert_paragraph_before, nu un singur șir: porțiunea înlocuită (eticheta care va deveni, la generare,
    numele/suma reală) moștenește formatarea PROPRIE a locului liber înlocuit (via `rpr_for`, de obicei
    _rpr_at pe paragraful original) — un „……” scris cu sublinire tot cu sublinire trebuie să iasă, chiar dacă
    restul frazei nu e; textul literal din jur rămâne pe formatarea dominantă a paragrafului (rpr_for=None)."""
    segments: list[tuple[str, object]] = []
    pos = len(text)
    for b in sorted(blanks, key=lambda x: x["start"], reverse=True):
        tag = tag_for(b)
        if not tag:
            continue
        s, e = b["start"] - offset, b["end"] - offset
        if e < pos:
            segments.append((text[e:pos], None))
        segments.append((tag, rpr_for(b) if rpr_for else None))
        pos = s
    if pos > 0:
        segments.append((text[:pos], None))
    segments.reverse()
    return segments


_CLAUSE_END_RE = re.compile(r"[.;]|-{4,}|,")
_DASH_FILL_AFTER_RE = re.compile(r"\s*-{4,}")    # liniuță de umplere imediat după punctul de final de propoziție


def _extend_clause_end(whole_text: str, end: int, limit: int, include_dash_fill: bool = False) -> int:
    """Extinde sfârșitul unei clauze (`end` = imediat după ultimul loc liber completat) până la o punctuație
    de final de propoziție („.”/„;”, inclusă), un bloc de liniuțe de umplere (vezi doc_filler._fix_dash_fill_tails
    — inclus DOAR dacă `include_dash_fill`, altfel exclus, tratat separat, la generare) sau o virgulă (exclusă)
    — oricare apare prima. Fără asta, restul propoziției pentru care CHIAR ACEA persoană e subiectul
    („ lei.”, „, cu puteri depline și cu o durată a mandatului până la …… ani.”) ar rămâne scris o singură
    dată, nu repetat pentru fiecare persoană (vezi _apply_inline_group/_apply_paragraph_group). Virgula
    oprește înadins — o virgulă imediat după ultimul câmp înseamnă de obicei o continuare colectivă, despre
    TOATE persoanele de-odată, nu despre cea curentă („……, CNP ……, aceștia fiind de acord.” — „aceștia” la
    plural se referă la toată lista, nu doar la ultima persoană din ea; rămâne sufix fix, o singură dată).

    `include_dash_fill=True` — folosit DOAR de _apply_paragraph_group, unde fiecare clauză e deja propriul ei
    paragraf (nu există risc de a înghiți clauza URMĂTOARE, `limit` fiind finalul ACELUIAȘI paragraf): un bloc
    de liniuțe de la finalul rândului aparține acelei persoane — „HOLHOS CADMIEL contribuie cu 2.000 lei.----”
    trebuie repetat identic pentru fiecare asociat, nu păstrat o singură dată ca separator de-o singură dată
    (asta rămâne comportamentul implicit la _apply_inline_group, unde mai multe clauze chiar pot împărți
    același paragraf, iar o liniuță găsită între ele chiar E un separator colectiv, nu al ultimei persoane).
    Cazul obișnuit e liniuța DUPĂ punct („…lei.----”): prima potrivire găsită e punctul, nu liniuța (apare
    mai devreme în text) — după ce punctul e inclus, se caută ȘI o liniuță imediat următoare (cu spații albe
    opționale între), ca amândouă să intre în clauza repetată, nu doar punctul.
    Nu trece de `limit` (începutul clauzei URMĂTOARE, sau finalul paragrafului) — altfel ar înghiți din
    clauza de după, când nu există nicio punctuație/liniuță/virgulă între ele."""
    m = _CLAUSE_END_RE.search(whole_text, end, limit)
    if not m:
        return limit
    if whole_text[m.start()] in ".;":
        pos = m.end()
        if include_dash_fill:
            dm = _DASH_FILL_AFTER_RE.match(whole_text, pos, limit)
            if dm:
                pos = dm.end()
        return pos
    if include_dash_fill and whole_text[m.start()] == "-":
        return m.end()
    return m.start()


def _split_group_paragraph(anchor: Paragraph, whole_text: str, template_span: tuple[int, int], group_end: int,
                           other_blanks: list[dict], template_blanks: list[dict], role: str,
                           choices: dict[int, str | None], prefix_end: int | None = None,
                           repeat_prefix: bool = False) -> None:
    """Un paragraf → prefix (text fix) / {{#ROL}} / clauza-șablon (etichete generice) / {{/ROL}} / sufix (text fix),
    fiecare ca paragraf separat, cu aceeași aliniere și font ca paragraful original (vezi _insert_paragraph_before
    și _dominant_rpr). Textul dintre sfârșitul clauzei-șablon și `group_end` (clauzele 2..N ale grupului,
    plus conectorul „si”/„iar” dintre ele, la varianta „inline”) se elimină — devine parte a blocului repetitiv,
    nu se mai scrie o singură dată. Blank-urile din afara grupului (prefix/sufix) primesc alegerea individuală a
    utilizatorului, exact ca la un loc liber obișnuit; cele lăsate „neschimbate” rămân „……” vizibil.

    `prefix_end` e limita prefixului (începutul primei clauze, în ordinea din text) — de obicei coincide cu
    începutul clauzei-șablon, dar nu întotdeauna: clauza-șablon poate fi alta decât prima (cea mai completă
    structural, ex. singura cu ADRESA), caz în care prefixul tot trebuie să se oprească înainte de PRIMA
    clauză, nu înainte de clauza-șablon — altfel clauza (clauzele) dinaintea ei ar rămâne text fix, needitat.

    `repeat_prefix=True` (folosit DOAR de _apply_paragraph_group, unde fiecare clauză e deja propriul ei
    paragraf — nu există concept de „prefix comun mai multor clauze” ca la _apply_inline_group): textul
    dinaintea primului câmp, dacă NU e doar un număr de ordine (vezi `prepend_index` mai jos), nu devine un
    paragraf fix de-o singură dată — se lipește la începutul FIECĂREI clauze repetate. „Asociatului …… îi
    revin …… părți sociale” trebuie să înceapă cu „Asociatului” la fiecare persoană generată, nu doar la
    prima — altfel „Asociatului” apărea o singură dată, ca titlu fals, deasupra întregii liste (bug real,
    găsit pe un document real)."""
    clause_start, clause_end = template_span
    if prefix_end is None:
        prefix_end = clause_start
    prefix_raw, clause_raw, suffix_raw = whole_text[:prefix_end], whole_text[clause_start:clause_end], whole_text[group_end:]
    # Pentru un rol necunoscut încă în tabelul de mai sus (comodant, reprezentant legal…) — plural aproximativ
    # (sufix „I”), suficient cât să fie un nume de bloc unic; nu ajunge niciodată vizibil ca text, doar ca
    # etichetă internă {{#ROL}}/{{/ROL}}, deci nu contează gramatical.
    role_plural = _ROLE_PLURAL.get(role, role + "I")

    # un prefix format DOAR dintr-un număr de ordine ("1. ") aparține de fapt clauzei, ca {{INDEX}} — nu rămâne
    # text fix, altfel fiecare persoană repetată ar purta numărul primei. Se marchează ACUM (pe textul original,
    # neschimbat) și se aplică mai jos, DUPĂ substituirea locurilor libere din clauză — altfel ar strica offset-urile.
    ordinal = _ORDINAL_RE.match(prefix_raw)
    prepend_index = ordinal is not None and ordinal.end() >= len(prefix_raw.rstrip())
    if prepend_index:
        prefix_raw = ""

    prefix_blanks = [b for b in other_blanks if b["end"] <= prefix_end]
    suffix_blanks = [b for b in other_blanks if b["start"] >= group_end]

    # fontul paragrafului original (Arial, Times New Roman, ce-o fi) — altfel paragrafele noi ar ieși cu
    # fontul implicit din stilul „Normal” al documentului, diferit de restul șablonului.
    rpr = _dominant_rpr(anchor)

    # Formatarea PROPRIE a fiecărui loc liber (ex. „……” subliniat) — nu doar cea dominantă a paragrafului —
    # vezi _local_replacements/_rpr_at; poziția e cea din `anchor`, ÎNCĂ neatins la acest punct.
    rpr_for = lambda b: _rpr_at(anchor, b["start"])   # noqa: E731

    repeat_this_prefix = repeat_prefix and not prepend_index and bool(prefix_raw.strip())
    if prefix_raw.strip() and not repeat_this_prefix:
        segs = _local_replacements(prefix_raw, prefix_blanks, 0, lambda b: choices.get(b["id"]), rpr_for)
        _insert_paragraph_before(anchor, segs, rpr)
    _insert_paragraph_before(anchor, f"{{{{#{role_plural}}}}}", rpr)
    clause_segs = _local_replacements(clause_raw, template_blanks, clause_start, lambda b: generic_person_tag(b["field"]), rpr_for)
    if prepend_index:
        if clause_segs:
            clause_segs[0] = (clause_segs[0][0].lstrip(), clause_segs[0][1])
        clause_segs.insert(0, ("{{INDEX}}. ", None))
    elif repeat_this_prefix:
        prefix_segs = _local_replacements(prefix_raw, prefix_blanks, 0, lambda b: choices.get(b["id"]), rpr_for)
        clause_segs = prefix_segs + clause_segs
    _insert_paragraph_before(anchor, clause_segs, rpr)
    _insert_paragraph_before(anchor, f"{{{{/{role_plural}}}}}", rpr)
    if suffix_raw.strip():
        segs = _local_replacements(suffix_raw, suffix_blanks, group_end, lambda b: choices.get(b["id"]), rpr_for)
        _insert_paragraph_before(anchor, segs, rpr)
    anchor._p.getparent().remove(anchor._p)


def _apply_inline_group(pars: list[Paragraph], group: dict, blanks_by_id: dict[int, dict], choices: dict[int, str | None]) -> None:
    par = pars[group["paragraph"]]
    text = "".join(r.text for r in par.runs)
    template_blanks = [blanks_by_id[i] for i in group["template_blank_ids"]]
    all_group_blanks = sorted((blanks_by_id[i] for i in group["blank_ids"]), key=lambda b: b["start"])
    prefix_end = all_group_blanks[0]["start"]    # începutul PRIMEI clauze — nu neapărat clauza-șablon (vezi mai jos)
    # Sfârșitul ULTIMEI clauze (clauzele 2..N dispar odată cu ce e între ele) — extins la fel ca mai jos, până
    # la finalul propoziției ei („ lei.”), nu doar până la ultimul loc liber, altfel acel rest ar rămâne
    # dublat: o dată (greșit) ca literă fixă de sufix, o dată prin extinderea clauzei-șablon de mai jos.
    group_end = _extend_clause_end(text, all_group_blanks[-1]["end"], len(text))
    # Restul propoziției clauzei-șablon, DUPĂ ultimul ei loc liber completat („ lei.”, „, cu puteri depline…”)
    # — până la clauza URMĂTOARE (dacă templateul nu e ultima), nu doar până la ultimul câmp (vezi
    # _extend_clause_end) — altfel acel rest ar rămâne scris o singură dată (ca sufix al grupului), nu
    # repetat pentru fiecare persoană.
    template_last_idx = all_group_blanks.index(template_blanks[-1])
    next_clause_start = all_group_blanks[template_last_idx + 1]["start"] if template_last_idx + 1 < len(all_group_blanks) else group_end
    template_span = (template_blanks[0]["start"], _extend_clause_end(text, template_blanks[-1]["end"], next_clause_start))
    other_blanks = [b for b in blanks_by_id.values() if b["paragraph"] == group["paragraph"] and b["id"] not in group["blank_ids"]]
    _split_group_paragraph(par, text, template_span, group_end, other_blanks, template_blanks, group["role"], choices, prefix_end)


def _apply_paragraph_group(pars: list[Paragraph], group: dict, blanks_by_id: dict[int, dict], choices: dict[int, str | None]) -> None:
    template_blanks = sorted((blanks_by_id[i] for i in group["template_blank_ids"]), key=lambda b: b["start"])
    first_pi = group["paragraphs"][0]
    first_par = pars[first_pi]
    text = "".join(r.text for r in first_par.runs)
    # Ca la _apply_inline_group — restul propoziției de după ultimul loc liber (vezi _extend_clause_end),
    # mărginit aici doar de finalul paragrafului (fiecare clauză e deja propriul ei paragraf). Excepție:
    # activitățile CAEN secundare (role="CAEN") sunt DOAR câte un cod pe rând, fără propoziție proprie —
    # punctuația de final aparține frazei introductive de dinaintea listei ("...activități:"), nu fiecărui
    # cod în parte; extinderea ar repeta-o greșit după fiecare cod. Pentru restul rolurilor, `include_dash_fill`
    # e pornit: liniuța de umplere de la finalul rândului (vezi _extend_clause_end) aparține clauzei ei —
    # trebuie să apară la fiecare persoană generată, nu doar o dată, ca separator fals după toată lista.
    is_caen = group["role"] == "CAEN"
    clause_end = template_blanks[-1]["end"] if is_caen else \
        _extend_clause_end(text, template_blanks[-1]["end"], len(text), include_dash_fill=True)
    template_span = (template_blanks[0]["start"], clause_end)
    other_blanks = [b for b in blanks_by_id.values() if b["paragraph"] == first_pi and b["id"] not in group["blank_ids"]]
    # repeat_prefix=True: fiecare paragraf al grupului E o clauză întreagă — orice text dinaintea primului
    # câmp (ex. „Asociatului”) aparține clauzei, nu e un titlu comun de-o singură dată (vezi mai sus).
    _split_group_paragraph(first_par, text, template_span, template_span[1], other_blanks, template_blanks,
                           group["role"], choices, repeat_prefix=True)
    for pi in group["paragraphs"][1:]:
        p = pars[pi]
        p._p.getparent().remove(p._p)


def apply(docx_bytes: bytes, choices: dict[int, str | None], groups: dict[int, str] | None = None) -> bytes:
    """Înlocuiește locurile libere cu etichetele alese ({id: etichetă}); id-urile lipsă sau None rămân neschimbate.

    `groups` ({group_id: "repeat"|"fixed"}) decide, pentru fiecare grup găsit de `detect_groups`, dacă persoanele
    lui devin un bloc repetitiv (scalează la orice număr) sau rămân poziții fixe (tratate ca locuri libere
    obișnuite, prin `choices`)."""
    doc = Document(io.BytesIO(docx_bytes))
    found = analyze(docx_bytes)
    blanks_by_id = {b["id"]: b for b in found}
    detected = detect_groups(found, docx_bytes)
    accepted = [g for g in detected if (groups or {}).get(g["id"]) == "repeat"]
    grouped_paragraphs = {g["paragraph"] for g in accepted if g["kind"] == "inline"} | \
        {pi for g in accepted if g["kind"] == "paragraph" for pi in g["paragraphs"]}

    by_par: dict[int, list[dict]] = {}
    for f in found:
        if f["paragraph"] in grouped_paragraphs:
            continue     # paragrafele unui grup acceptat sunt tratate integral de _apply_*_group, mai jos
        tag = choices.get(f["id"])
        if tag:
            by_par.setdefault(f["paragraph"], []).append({**f, "tag": tag})
    pars = _paragraphs(doc)
    for pi, items in by_par.items():
        for it in sorted(items, key=lambda x: x["start"], reverse=True):
            _replace_span(pars[pi], it["start"], it["end"], it["tag"])

    for g in accepted:
        if g["kind"] == "inline":
            _apply_inline_group(pars, g, blanks_by_id, choices)
        else:
            _apply_paragraph_group(pars, g, blanks_by_id, choices)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


_TAG_OK = re.compile(r"^(?:\{\{[A-Z0-9_]+\}\}\s?){1,3}$")


def valid_tag(tag) -> bool:
    """Doar etichete {{CÂMP}} (cel mult trei, ex. nume + prenume); nimic altceva nu ajunge în document."""
    return isinstance(tag, str) and 0 < len(tag) <= 120 and bool(_TAG_OK.match(tag.strip()))
